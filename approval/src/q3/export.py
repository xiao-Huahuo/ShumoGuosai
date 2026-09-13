"""3.21/3.27：真实费用表、原模板回填与逐值回读；禁止部分结果冒充完整结果。"""
from copy import copy
from pathlib import Path
import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from .config import ATTACHMENTS, Config, write_csv, write_json, sync_directory
from .rolling import validate_frame


def intervals() -> list[str]:
    return [f'{t//6:02d}:{t%6*10:02d}-{(t+1)//6:02d}:{(t+1)%6*10:02d}' for t in range(144)]


def emergency_events(frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    current = None
    for row in frame.itertuples():
        end = pd.Timestamp(row.timestamp); start = end-pd.Timedelta(minutes=10)
        if row.emergency > 0:
            if current is None or current['end'] != start:
                if current is not None:
                    rows.append(current)
                current = {'date': str(start.date()), 'start': start, 'end': end, 'energy_kwh': 0., 'cost_yuan': 0.}
            current['end'] = end
            current['energy_kwh'] += row.emergency
            current['cost_yuan'] += row.emergency_cost
        elif current is not None:
            rows.append(current); current = None
    if current is not None:
        rows.append(current)
    for row in rows:
        offset = (row['end'].normalize()-row['start'].normalize()).days
        row['interval'] = f'{row["start"]:%H:%M}-{row["end"]:%H:%M}'+(f'+{offset}日' if offset else '')
    return pd.DataFrame(rows, columns=['date', 'start', 'end', 'interval', 'energy_kwh', 'cost_yuan'])


def summaries(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    daily, blocks = [], []
    for date, day in frame.groupby('date', sort=False):
        events = emergency_events(day)
        daily.append({'date': date, **{key: float(day[key].sum()) for key in
            ('g0', 'grid', 'planned_cost', 'adjustment_cost', 'emergency_cost', 'total_cost', 'emergency', 'charge', 'discharge', 'spill')},
            'emergency_events': len(events), 'emergency_slots': int((day.emergency > 0).sum()),
            'throughput': float(day.charge.sum()+day.discharge.sum()),
            'soc_start': float(day.initial_soc.iloc[0]), 'soc_end': float(day.soc.iloc[-1])})
        for block in range(6):
            part = day.iloc[block*24:(block+1)*24]
            blocks.append({'date': date, 'interval': f'{block*4:02d}:00-{(block+1)*4:02d}:00',
                'charge': float(part.charge.sum()), 'discharge': float(part.discharge.sum()),
                'soc_start': float(day.initial_soc.iloc[0]), 'soc_end': float(day.soc.iloc[-1])})
    return pd.DataFrame(daily), pd.DataFrame(blocks)


def write_tables(frame: pd.DataFrame, output: Path) -> None:
    daily, blocks = summaries(frame)
    write_csv(output/'daily.csv', daily)
    write_csv(output/'storage_blocks.csv', blocks)
    write_csv(output/'emergency_events.csv', emergency_events(frame))
    dates = ['2025-03-20', '2025-06-21', '2025-09-23', '2025-12-21']
    chosen = frame[frame.date.astype(str).isin(dates)]
    table1 = chosen[chosen.slot.isin([60, 72, 84, 96, 108, 120])].copy()
    table1['interval'] = [intervals()[int(t)] for t in table1.slot]
    write_csv(output/'paper_table1.csv', table1[['date', 'interval', 'g0', 'grid', 'planned_cost', 'adjustment_cost', 'emergency_cost', 'total_cost']]
              .merge(daily[['date', 'grid', 'total_cost']], on='date', suffixes=('', '_daily')))
    write_csv(output/'paper_table2.csv', blocks[blocks.date.astype(str).isin(dates)])
    events = [emergency_events(part) for _, part in chosen.groupby('date')]
    write_csv(output/'paper_table3.csv', pd.concat(events, ignore_index=True) if events else emergency_events(chosen))


def export_workbook(frame: pd.DataFrame, output: Path, config: Config, *, smoke: bool = False) -> dict:
    validate_frame(frame, config, full=not smoke)
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(ATTACHMENTS/'附件5/result3.xlsx')
    expected = ['计划购电量', '调整购电量', '充放电量', '紧急购电量']
    if workbook.sheetnames != expected:
        raise ValueError('原模板工作表结构发生改变')
    originals = {s.title: [c.value for c in s[1]] for s in list(workbook)[:2]}
    styles = {s.title: [copy(c._style) for c in s[2]] for s in workbook}
    for sheet in workbook:
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        sheet.delete_rows(2, sheet.max_row)
    plans, adjusted, storage, emergencies = [workbook[name] for name in expected]
    for sheet in (plans, adjusted):
        for col, label in enumerate(intervals(), 2):
            sheet.cell(1, col, label)
        sheet.cell(1, 2).comment = Comment('原模板时段整体偏移10分钟。按附件右端点统一回填00:00—24:00共144段，原标签记录于template_audit.json。', 'Q3')
    adjusted.cell(1, 2).comment = Comment('调整购电量为最终生效绝对量；0—6点保持原计划。差额和逐节点proxy分别存CSV及节点审计，不混入真实费用。原时段标签错位已修正。', 'Q3')
    for i, (date, day) in enumerate(frame.groupby('date', sort=False), 2):
        for sheet, key, cost in ((plans, 'g0', 'planned_cost'), (adjusted, 'grid', 'adjustment_cost')):
            sheet.cell(i, 1, pd.Timestamp(date).to_pydatetime())
            for col, value in enumerate(day[key], 2):
                sheet.cell(i, col, float(value))
            # 原比赛模板为仿真结果的数值输入区，沿用模板及Q2惯例，非用户可编辑计算模型。
            sheet.cell(i, 146, float(day[key].sum()))
            sheet.cell(i, 147, float(day[cost].sum()))
    adjusted.cell(1, 147).comment = Comment('调整相关净费用=Σ[1.5p增量−0.5p取消量]；负数表示相对原计划费用减少。总真实费用另存daily.csv。', 'Q3')
    _, blocks = summaries(frame)
    for i, row in enumerate(blocks.itertuples(index=False), 2):
        block = (i-2)%6
        values = (pd.Timestamp(row.date).to_pydatetime() if block == 0 else None, row.interval, row.charge, row.discharge,
                  '0:00' if block == 0 else '24:00' if block == 1 else None,
                  row.soc_start if block == 0 else row.soc_end if block == 1 else None)
        for col, value in enumerate(values, 1):
            storage.cell(i, col, value)
    events = emergency_events(frame)
    for i, row in enumerate(events.itertuples(index=False), 2):
        for col, value in enumerate((pd.Timestamp(row.date).to_pydatetime(), row.interval, row.energy_kwh), 1):
            emergencies.cell(i, col, value)
    for sheet in workbook:
        sheet.freeze_panes = 'B2' if sheet in (plans, adjusted) else 'A2'
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell._style = copy(styles[sheet.title][cell.column-1])
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '0.000000'
    filename = 'smoke_result3.xlsx' if smoke else 'result3.xlsx'
    temporary = output/(filename+'.pending.xlsx')
    workbook.save(temporary); workbook.close()
    check = load_workbook(temporary, read_only=True, data_only=True)
    try:
        for sheet, key, cost in (('计划购电量', 'g0', 'planned_cost'), ('调整购电量', 'grid', 'adjustment_cost')):
            saved = list(check[sheet].values)[1:]
            np.testing.assert_allclose(np.array([r[1:145] for r in saved], dtype=float).ravel(), frame[key], atol=1e-8)
            np.testing.assert_allclose(np.array([r[145:147] for r in saved], dtype=float), frame.groupby('date', sort=False)[[key, cost]].sum(), atol=1e-8)
        saved = list(check['充放电量'].values)[1:]
        np.testing.assert_allclose(np.array([r[2:4] for r in saved], dtype=float), blocks[['charge', 'discharge']], atol=1e-8)
        for i, row in enumerate(saved):
            if i%6 in (0, 1):
                np.testing.assert_allclose(row[5], blocks.iloc[i]['soc_start' if i%6 == 0 else 'soc_end'], atol=1e-8)
        saved_events = list(check['紧急购电量'].values)[1:]
        np.testing.assert_allclose(np.asarray([r[2] for r in saved_events], dtype=float), events.energy_kwh.to_numpy(dtype=float), atol=1e-8)
    finally:
        check.close()
    with temporary.open('rb') as stream:
        os.fsync(stream.fileno())
    temporary.replace(output/filename)
    sync_directory(output)
    audit = {'smoke_only': smoke, 'days': int(frame.date.nunique()), 'roundtrip_all_sheets': True,
             'original_headers': originals, 'intervals': intervals(), 'adjusted_quantity': 'final_effective_absolute',
             'storage_from': 'actual_replay', 'full_period': not smoke}
    write_json(output/'template_audit.json', audit)
    return audit
