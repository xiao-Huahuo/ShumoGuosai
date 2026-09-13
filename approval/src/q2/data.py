"""2.1.1–2.1.5：保留原值，核验宽表，建立右端点连续功率序列。"""

import datetime as dt
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
SLOTS = 144
DT_HOURS = 1 / 6
START = pd.Timestamp("2025-01-01")
END = pd.Timestamp("2025-12-31")
FORECAST_START = pd.Timestamp("2025-02-01")
SERIES = ("load_kw", "pv_kw", "net_load_kw")
CALENDAR = ("date", "interval", "timestamp", "slot", "weekday", "weekday_name", "month")
SHEETS = {"load_kw": "小区负载", "pv_kw": "光伏发电实际功率"}


def sha256(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def write_json(path, value):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def write_csv(path, frame):
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    frame.to_csv(path, index=False, encoding="utf-8", float_format="%.17g")


def endpoint_minutes(value):
    """附件时间的唯一含义为区间右端点；不把裸 0:00 猜成 24:00。"""
    if isinstance(value, dt.time):
        if value.second or value.microsecond:
            raise ValueError("时间存在秒数")
        minutes = value.hour * 60 + value.minute
    elif isinstance(value, str):
        if value.strip() in ("0:00+1", "00:00+1", "24:00"):
            return 1440
        hour, minute = map(int, value.strip().split(":"))
        if not (0 <= hour < 24 and 0 <= minute < 60):
            raise ValueError("非法时刻")
        minutes = hour * 60 + minute
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        if not np.isfinite(value) or not np.isclose(value * 1440, round(value * 1440), atol=1e-8, rtol=0):
            raise ValueError("非法 Excel 时间日分数")
        minutes = round(value * 1440)
    else:
        raise ValueError("无法识别的时间格式")
    if minutes not in range(10, 1441, 10):
        raise ValueError("右端点必须为 00:10 至 24:00 的 10 分钟倍数")
    return minutes


def calendar_for_day(date):
    date = pd.Timestamp(date).normalize()
    minute = np.arange(1, SLOTS + 1) * 10
    def clock(m):
        return f"{m // 60:02d}:{m % 60:02d}"
    return pd.DataFrame({
        "date": date, "interval": [f"{clock(m - 10)}-{clock(m)}" for m in minute],
        "timestamp": date + pd.to_timedelta(minute, unit="min"), "slot": np.arange(1, SLOTS + 1),
        "weekday": date.dayofweek + 1, "weekday_name": "星期" + "一二三四五六日"[date.dayofweek],
        "month": date.month,
    })


def read_attachment(path, qc_path):
    """硬错误记录后拒绝继续；不存在统计截尾、插值或逐点改值。"""
    checks, arrays = [], {}
    def check(name, passed, detail):
        checks.append({"check": name, "status": "pass" if passed else "fail", "detail": str(detail)})
    workbook = load_workbook(path, read_only=True, data_only=True)
    dates = pd.date_range(START, END, freq="D")
    try:
        check("required_sheets", set(SHEETS.values()).issubset(workbook.sheetnames), workbook.sheetnames)
        for variable, sheet_name in SHEETS.items():
            if sheet_name not in workbook.sheetnames:
                continue
            rows = list(workbook[sheet_name].iter_rows(values_only=True))
            check(f"{variable}:shape", len(rows) == 366 and all(len(r) == 145 for r in rows),
                  f"rows={len(rows)}, columns={len(rows[0]) if rows else 0}")
            if len(rows) != 366 or any(len(r) != 145 for r in rows):
                continue
            try:
                minutes = [endpoint_minutes(v) for v in rows[0][1:]]
                check(f"{variable}:144_unique_endpoints", sorted(minutes) == list(range(10, 1441, 10)), minutes)
                observed_dates = pd.DatetimeIndex(pd.to_datetime([r[0] for r in rows[1:]], errors="raise"))
                check(f"{variable}:365_unique_dates", observed_dates.is_unique and observed_dates.sort_values().equals(dates),
                      f"unique={observed_dates.nunique()}, first={observed_dates.min()}, last={observed_dates.max()}")
                values = np.array([r[1:] for r in rows[1:]], dtype=float)
                booleans = sum(isinstance(v, bool) for r in rows[1:] for v in r[1:])
                check(f"{variable}:numeric_not_boolean", booleans == 0, f"boolean_count={booleans}")
                check(f"{variable}:finite_no_missing", np.isfinite(values).all(),
                      f"nonfinite_count={np.count_nonzero(~np.isfinite(values))}")
                check(f"{variable}:nonnegative", (values >= 0).all(), f"negative_count={np.count_nonzero(values < 0)}")
                arrays[variable] = values[np.argsort(observed_dates)][:, np.argsort(minutes)]
            except (ValueError, TypeError, OverflowError) as error:
                check(f"{variable}:format", False, error)
    finally:
        workbook.close()
    write_csv(qc_path, pd.DataFrame(checks))
    if any(c["status"] == "fail" for c in checks):
        raise ValueError(f"附件核验失败，原始值未修改；见 {qc_path}")
    frame = pd.concat([calendar_for_day(d) for d in dates], ignore_index=True)
    for variable in SHEETS:
        frame[variable] = arrays[variable].ravel()
    frame["net_load_kw"] = frame.load_kw - frame.pv_kw
    check("52560_continuous_unique_timestamps", len(frame) == 52560 and frame.timestamp.is_unique and
          frame.timestamp.diff().dropna().eq(pd.Timedelta(minutes=10)).all(),
          f"{frame.timestamp.iloc[0]} to {frame.timestamp.iloc[-1]}")
    check("144_periods_each_source_date", frame.groupby("date").size().eq(SLOTS).all(), "按原始日期归属，非右端点自然日分组")
    check("zero_pv_preserved", int((frame.pv_kw == 0).sum()) == int((arrays["pv_kw"] == 0).sum()),
          f"zero_count={int((frame.pv_kw == 0).sum())}; 零值合法，未指定或推断固定夜间时段")
    write_csv(qc_path, pd.DataFrame(checks))
    if any(c["status"] == "fail" for c in checks):
        raise ValueError("长表连续性检查失败")
    return frame, pd.DataFrame(checks)


def history_before(frame, origin):
    """用户确认：严格小于；上一日末段即使已结束也排除。"""
    origin = pd.Timestamp(origin)
    if origin != origin.normalize():
        raise ValueError("预测时点必须为 0:00")
    return frame.loc[frame.timestamp < origin].copy(deep=True)


def jump_context(frame):
    """列出所有局部转折点，按偏离相邻值程度排序；仅供人工复核，不判异常。"""
    rows = []
    for variable in SHEETS:
        x = frame[variable]
        previous, following = x.shift(1), x.shift(-1)
        reversal = ((x - previous) * (following - x) < 0)
        view = frame.loc[reversal, ["date", "timestamp", "slot", "weekday"]].copy()
        view["series"] = variable
        view["value_kw"] = x[reversal]
        view["previous_kw"] = previous[reversal]
        view["following_kw"] = following[reversal]
        view["neighbor_deviation_kw"] = (x - (previous + following) / 2)[reversal]
        view["previous_day_kw"] = x.shift(144)[reversal]
        view["previous_week_kw"] = x.shift(1008)[reversal]
        groups = frame.groupby(["weekday", "slot"])[variable]
        for method in ("min", "median", "max", "count"):
            historic = groups.transform(lambda v: getattr(v.shift(1).expanding(), method)())
            view[f"past_same_weekday_{method}"] = historic[reversal]
        view["action"] = "retain_review_only"
        view = view.loc[view.neighbor_deviation_kw.abs().sort_values(ascending=False).index]
        view["review_rank"] = np.arange(1, len(view) + 1)
        rows.append(view)
    return pd.concat(rows, ignore_index=True)
