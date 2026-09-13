"""全年完成后的独立验收：逐节点/物理/SHA/七个Excel工作表与逐格日期和SOC。"""
from pathlib import Path
import sys,json,hashlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[3]
sys.path.insert(0,str(ROOT))


def check_mode(mode):
    short='q4'+mode[-1]
    version='v2' if mode=='4-2' else 'v3'
    folder=ROOT/f'outputs/q4/raw/rescue_lp_{short}_20260913_{version}'
    state=json.loads((folder/'state.json').read_text(encoding='utf-8'))
    assert len(state['days'])==334
    assert [a['day'] for a in state['days']]==list(range(31,365))
    dispatch=pd.read_csv(folder/'dispatch.csv',float_precision='round_trip')
    expected_dates=pd.date_range('2025-02-01',periods=334)
    expected_times=pd.date_range('2025-02-01 00:10',periods=334*144,freq='10min')
    np.testing.assert_array_equal(pd.to_datetime(dispatch.timestamp),expected_times)
    np.testing.assert_array_equal(dispatch.slot,np.tile(np.arange(144),334))
    assert len(dispatch)==48096
    c=dispatch.charge.to_numpy();r=dispatch.discharge.to_numpy()
    soc=6000+np.cumsum(.9*c-r/.9)
    soc_error=float(np.max(abs(soc-dispatch.soc)))
    assert soc_error<=1e-6
    assert dispatch.soc.min()>=1200-1e-6 and dispatch.soc.max()<=10800+1e-6
    assert np.max(c+r)<=5000/6+1e-6
    assert np.max(np.minimum(c,r))<=1e-6
    balance=float(abs(dispatch.called+dispatch.discharge+dispatch.emergency-dispatch.net_kwh-dispatch.charge-dispatch.spill).max())
    assert balance<=1e-6
    difference=dispatch.grid-dispatch.g0
    adjustment=np.where(difference>=0,1.5*difference,.5*difference) if mode=='4-3' else np.zeros(len(dispatch))
    np.testing.assert_allclose(dispatch.total_cost,dispatch.price*(dispatch.g0+adjustment+5*dispatch.emergency),atol=1e-6,rtol=0)
    audits=[]
    for index,date in enumerate(expected_dates):
        day=dispatch.iloc[index*144:(index+1)*144]
        receipt=state['days'][index]
        for key,sub,suffix in [('dispatch','days','.csv'),('audit','audit','.json')]:
            path=folder/sub/(str(date.date())+suffix)
            assert hashlib.sha256(path.read_bytes()).hexdigest()==receipt['hashes'][key]
        audit=json.loads((folder/'audit'/f'{date.date()}.json').read_text(encoding='utf-8'))
        nodes=[audit] if mode=='4-2' else audit['nodes']
        if mode=='4-3':assert [n['hour'] for n in nodes]==[0,6,12,18]
        for node in nodes:
            hour=0 if mode=='4-2' else node['hour'];count=node['committed_slots'];start=hour*6
            part=day.iloc[start:start+count]
            for key in ['grid','charge','discharge']:
                np.testing.assert_allclose(part[key],node[key][:count],atol=1e-6,rtol=0)
            initial=6000 if index==0 and hour==0 else float(dispatch.soc.iloc[index*144+start-1])
            assert abs(node['initial_soc']-initial)<=1e-6
            solver=node['solver'];scenario=node['scenarios'];audits.append(solver)
            assert solver['status']=='Optimal' and solver['reliable'] and solver['valid']
            assert solver['solution_source']=='optimal_solver' and solver['integer_variables']==0
            assert solver['storage_schedule']=='node_shared_fixed_committed_block'
            assert solver['dro_coupling_constraints']==scenario['S']**2
            assert scenario['window']==56 and scenario['joint'] and scenario['bootstrap_repetitions']==100
            assert scenario['S']<=20
            assert scenario['tail_singletons_verified']
            np.testing.assert_allclose(scenario['tail_prior_mass'],scenario['tail_weights'],atol=1e-12,rtol=0)
            assert all(j*144+hour*6+scenario['length']<=receipt['day']*144+hour*6 for j in scenario['history_days'])
            assert all(j>=receipt['day']-56 for j in scenario['history_days'])
    assert len(audits)==(334 if mode=='4-2' else 1336)
    workbook_path=folder/f'result{mode}.xlsx'
    wb=load_workbook(workbook_path,read_only=True,data_only=False)
    expected_sheets=['计划购电量']+(['调整购电量'] if mode=='4-3' else [])+['充放电量','紧急购电量']
    assert wb.sheetnames==expected_sheets
    maximum_error=0.
    def compare(a,b):
        nonlocal maximum_error
        a=np.asarray(a,dtype=float);b=np.asarray(b,dtype=float)
        if a.size:maximum_error=max(maximum_error,float(np.max(abs(a-b))))
        np.testing.assert_allclose(a,b,atol=1e-7,rtol=0)
    for sheet,key,cost in [('计划购电量','grid' if mode=='4-2' else 'g0','planned_cost')]+([('调整购电量','grid','adjustment_cost')] if mode=='4-3' else []):
        values=list(wb[sheet].values)
        assert len(values)==335
        for i,row in enumerate(values[1:]):
            assert pd.Timestamp(row[0])==expected_dates[i]
            day=dispatch.iloc[i*144:(i+1)*144]
            compare(row[1:145],day[key])
            compare(row[145:147],[day[key].sum(),day[cost].sum()])
        for slot,label in enumerate(values[0][1:145]):
            start=slot*10;end=(slot+1)*10
            assert label==f'{start//60:02d}:{start%60:02d}-{end//60:02d}:{end%60:02d}'
    storage=list(wb['充放电量'].values)[1:]
    assert len(storage)==334*6
    for i,date in enumerate(expected_dates):
        day=dispatch.iloc[i*144:(i+1)*144]
        block=storage[i*6:(i+1)*6]
        assert pd.Timestamp(block[0][0])==date
        compare([block[0][5],block[1][5]],[day.initial_soc.iloc[0],day.soc.iloc[-1]])
        assert block[0][4]=='0:00' and block[1][4]=='24:00'
        for k,row in enumerate(block):
            assert row[1]==f'{k*4:02d}:00-{(k+1)*4:02d}:00'
            part=day.iloc[k*24:(k+1)*24]
            compare(row[2:4],[part.charge.sum(),part.discharge.sum()])
    # 独立分组逐日连续正紧急购电，避免使用被测export的event实现。
    expected=[]
    for date,day in dispatch.groupby('date',sort=False):
        values=day.emergency.to_numpy();start=None
        for i in range(145):
            positive=i<144 and values[i]>0
            if positive and start is None:start=i
            if not positive and start is not None:
                label=f'{start//6:02d}:{start%6*10:02d}-{i//6:02d}:{i%6*10:02d}'
                # 导出当前采用strftime午夜00:00；保持题目右端点与实际导出规则。
                if i==144:label=label[:-5]+'00:00'
                expected.append((pd.Timestamp(date),label,float(values[start:i].sum())))
                start=None
    saved=list(wb['紧急购电量'].values)[1:]
    assert len(saved)==len(expected)
    for row,wanted in zip(saved,expected):
        assert pd.Timestamp(row[0])==wanted[0] and row[1]==wanted[1]
        compare([row[2]],[wanted[2]])
    error_values={'#REF!','#DIV/0!','#VALUE!','#NAME?','#N/A','#NUM!','#NULL!','#SPILL!','#CALC!'}
    for sheet in wb:
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type not in ('e','f')
                assert not isinstance(cell.value,str) or cell.value not in error_values
    wb.close()
    gates=json.loads((folder/'production_gates.json').read_text())
    assert gates['passed']
    assert gates['metrics']['C_realized']<14836427.918189546*1.15
    # 完成态收据显式汇总已验证的实际版本，避免仅凭路径追溯。
    from src.q4.config import write_json
    source=json.loads((folder/'source_manifest.json').read_text(encoding='utf-8'))
    solver_versions=sorted({a['solver_version'] for a in audits})
    assert len(solver_versions)==1
    state.update(solver_version=solver_versions[0], source_hashes=source['production_code_sha256'],
                 input_hashes={**source['q3_provenance']['files'], source['attachment4']:source['attachment4_sha256']},
                 q2_frozen_source={'path':source['q2_run'],'dispatch_sha256':source['q2_frozen_dispatch_sha256']},
                 price_forecast_version=source['price_code_sha256'],
                 price_forecasts_array_sha256=source['price_forecasts_array_sha256'],
                 W=state['config']['residual_window'], S=state['config']['scenarios'],
                 bootstrap_repetitions=state['config']['bootstrap_repetitions'])
    write_json(folder/'state.json',state)
    return dict(mode=mode,days=334,rows=48096,certified_nodes=len(audits),daily_hashes=668,
                integer_variables=0,max_balance_error=balance,max_soc_chain_error=soc_error,
                excel_all_cells_checked=True,excel_max_numeric_error=maximum_error,
                emergency_events=len(expected),formula_error_cells=0,
                also_passes_review_baseline_15pct=True,workbook=str(workbook_path),
                workbook_sha256=hashlib.sha256(workbook_path.read_bytes()).hexdigest())


if __name__=='__main__':
    requested=sys.argv[1:] or ['4-2','4-3']
    for mode in requested:
        result=check_mode(mode)
        (ROOT/'docs/4/rescue_lp'/f'full_acceptance_{mode}.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
        print(json.dumps(result,ensure_ascii=False),flush=True)
