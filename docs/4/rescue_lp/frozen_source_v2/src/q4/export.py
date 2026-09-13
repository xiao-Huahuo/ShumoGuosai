"""7.8：真实费用表和result4-2/result4-3模板回填、回读验收。"""
from copy import copy
from pathlib import Path
import os
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from .config import ATTACHMENTS, Config, write_csv, write_json
from .rolling import validate_frame


def intervals() -> list[str]:
    return [f"{slot // 6:02d}:{slot % 6 * 10:02d}-{(slot + 1) // 6:02d}:{(slot + 1) % 6 * 10:02d}"
            for slot in range(144)]


def emergency_events(frame: pd.DataFrame) -> pd.DataFrame:
    rows, current = [], None
    for row in frame.itertuples():
        end = pd.Timestamp(row.timestamp)
        start = end - pd.Timedelta(minutes=10)
        if row.emergency > 0:
            if current is None or current["end"] != start or current["date"] != str(start.date()):
                if current is not None:
                    rows.append(current)
                current = {"date": str(start.date()), "start": start, "end": end,
                           "energy_kwh": 0.0, "cost_yuan": 0.0}
            current["end"] = end
            current["energy_kwh"] += row.emergency
            current["cost_yuan"] += row.emergency_cost
        elif current is not None:
            rows.append(current)
            current = None
    if current is not None:
        rows.append(current)
    for row in rows:
        row["interval"] = f"{row['start']:%H:%M}-{row['end']:%H:%M}"
    return pd.DataFrame(rows, columns=["date", "start", "end", "interval", "energy_kwh", "cost_yuan"])


def summaries(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    daily, blocks = [], []
    for date, day in frame.groupby("date", sort=False):
        daily.append({"date": date, **{key: float(day[key].sum()) for key in
                     ("g0", "grid", "planned_cost", "adjustment_cost", "emergency_cost", "total_cost",
                      "emergency", "charge", "discharge", "spill", "unused")},
                      "cost_q95_component": float(day.total_cost.sum()),
                      "soc_start": float(day.initial_soc.iloc[0]), "soc_end": float(day.soc.iloc[-1])})
        for block in range(6):
            part = day.iloc[block * 24:(block + 1) * 24]
            blocks.append({"date": date, "interval": f"{block * 4:02d}:00-{(block + 1) * 4:02d}:00",
                           "charge": float(part.charge.sum()), "discharge": float(part.discharge.sum()),
                           "soc_start": float(day.initial_soc.iloc[0]), "soc_end": float(day.soc.iloc[-1])})
    return pd.DataFrame(daily), pd.DataFrame(blocks)


def write_tables(frame: pd.DataFrame, output: Path) -> None:
    daily, blocks = summaries(frame)
    write_csv(output / "daily.csv", daily)
    write_csv(output / "storage_blocks.csv", blocks)
    write_csv(output / "emergency_events.csv", emergency_events(frame))
    metrics = {
        "C_realized": float(frame.total_cost.sum()),
        "E_emergency": float(frame.emergency.sum()),
        "Q_0.95_daily_cost": float(daily.total_cost.quantile(0.95)),
        "adjustment_energy": float(abs(frame.grid - frame.g0).sum()),
        "price_MAE": float(abs(frame.price - frame.price_forecast).mean()),
        "price_RMSE": float(np.sqrt(np.mean((frame.price - frame.price_forecast) ** 2))),
    }
    write_json(output / "metrics.json", metrics)


def export_workbook(frame: pd.DataFrame, output: Path, mode: str, config: Config, *, smoke: bool = False) -> dict:
    validate_frame(frame, mode, config, full=not smoke)
    if not smoke:
        import hashlib
        import json
        gate_path = output / 'production_gates.json'
        if not gate_path.exists():
            raise ValueError('正式Excel必须先通过生产门禁')
        gates = json.loads(gate_path.read_text(encoding='utf-8'))
        if not gates['passed'] or gates['mode'] != mode or gates['dispatch_sha256'] != hashlib.sha256((output / 'dispatch.csv').read_bytes()).hexdigest():
            raise ValueError('生产门禁失败或验证后数据变动')
    expected = (["计划购电量", "充放电量", "紧急购电量"] if mode == "4-2" else
                ["计划购电量", "调整购电量", "充放电量", "紧急购电量"])
    template = ATTACHMENTS / "附件5" / f"result{mode}.xlsx"
    workbook = load_workbook(template)
    if workbook.sheetnames != expected:
        raise ValueError("问题四模板工作表结构发生变化")
    styles = {sheet.title: [copy(cell._style) for cell in sheet[2]] for sheet in workbook}
    for sheet in workbook:
        for merged in list(sheet.merged_cells.ranges):
            sheet.unmerge_cells(str(merged))
        sheet.delete_rows(2, sheet.max_row)
    quantity_sheets = [(workbook["计划购电量"], "grid" if mode == "4-2" else "g0", "planned_cost")]
    if mode == "4-3":
        quantity_sheets.append((workbook["调整购电量"], "grid", "adjustment_cost"))
    for sheet, _, _ in quantity_sheets:
        for column, label in enumerate(intervals(), 2):
            sheet.cell(1, column, label)
        sheet.cell(1, 2).comment = Comment("按10分钟右端点纠正为00:00—24:00共144段。", "Q4")
    for row_index, (date, day) in enumerate(frame.groupby("date", sort=False), 2):
        for sheet, key, cost in quantity_sheets:
            sheet.cell(row_index, 1, pd.Timestamp(date).to_pydatetime())
            for column, value in enumerate(day[key], 2):
                sheet.cell(row_index, column, float(value))
            sheet.cell(row_index, 146, float(day[key].sum()))
            sheet.cell(row_index, 147, float(day[cost].sum()))
    if mode == "4-3":
        workbook["调整购电量"].cell(1, 147).comment = Comment(
            "净调整费=Σp[1.5增购量−0.5减购量]；数量列为最终生效绝对购电量。", "Q4"
        )
    _, blocks = summaries(frame)
    storage = workbook["充放电量"]
    for row_index, row in enumerate(blocks.itertuples(index=False), 2):
        block = (row_index - 2) % 6
        values = (pd.Timestamp(row.date).to_pydatetime() if block == 0 else None, row.interval,
                  row.charge, row.discharge, "0:00" if block == 0 else "24:00" if block == 1 else None,
                  row.soc_start if block == 0 else row.soc_end if block == 1 else None)
        for column, value in enumerate(values, 1):
            storage.cell(row_index, column, value)
    events = emergency_events(frame)
    emergency = workbook["紧急购电量"]
    for row_index, row in enumerate(events.itertuples(index=False), 2):
        for column, value in enumerate((pd.Timestamp(row.date).to_pydatetime(), row.interval, row.energy_kwh), 1):
            emergency.cell(row_index, column, value)
    for sheet in workbook:
        sheet.freeze_panes = "B2" if sheet in [item[0] for item in quantity_sheets] else "A2"
        style = styles[sheet.title]
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.column <= len(style):
                    cell._style = copy(style[cell.column - 1])
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
    output.mkdir(parents=True, exist_ok=True)
    filename = f"smoke_result{mode}.xlsx" if smoke else f"result{mode}.xlsx"
    temporary = output / (filename + ".pending.xlsx")
    workbook.save(temporary)
    workbook.close()
    check = load_workbook(temporary, read_only=True, data_only=True)
    try:
        for sheet, key, cost in [(item[0].title, item[1], item[2]) for item in quantity_sheets]:
            saved = list(check[sheet].values)[1:]
            np.testing.assert_allclose(np.asarray([row[1:145] for row in saved], float).ravel(), frame[key], atol=1e-8)
            expected_totals = frame.groupby("date", sort=False)[[key, cost]].sum().to_numpy()
            np.testing.assert_allclose(np.asarray([row[145:147] for row in saved], float), expected_totals, atol=1e-8)
        saved_blocks = list(check["充放电量"].values)[1:]
        np.testing.assert_allclose(np.asarray([row[2:4] for row in saved_blocks], float),
                                   blocks[["charge", "discharge"]], atol=1e-8)
        saved_events = list(check["紧急购电量"].values)[1:]
        np.testing.assert_allclose(np.asarray([row[2] for row in saved_events], float),
                                   events.energy_kwh.to_numpy(float), atol=1e-8)
    finally:
        check.close()
    with temporary.open("r+b") as stream:
        os.fsync(stream.fileno())
    temporary.replace(output / filename)
    audit = {"mode": mode, "smoke_only": smoke, "days": int(frame.date.nunique()),
             "roundtrip_all_sheets": True, "adjusted_quantity": "final_effective_absolute"}
    write_json(output / f"template_audit_{mode}.json", audit)
    if not smoke:
        write_json(output / "active.json", {"status": "complete", "mode": mode, "completed_days": 334, "workbook": filename})
    return audit
