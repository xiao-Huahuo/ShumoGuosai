import tempfile
import unittest
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from src.q3.physics import Policy
from .config import Config, ROOT
from .data import read_inputs, read_prices
from .diagnostics import experiment_variants
from .export import export_workbook
from .optimization import solve, solve_dual
from .physics import adjustment_quantity, replay
from .price import feature_matrix
from .rolling import _cold_start, validate_frame
from .scenarios import Scenarios, _reduce, construct
from .traceability import generate


class Q4Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.data = read_inputs(ROOT / "inputs/q4/processed")

    def test_actual_price_shape_alignment_and_reported_statistics(self):
        prices = read_prices(ROOT / "docs/CUMCM2026Problems/C题/附件/附件4.xlsx")
        self.assertEqual(prices.shape, (365, 144))
        self.assertAlmostEqual(float(prices.min()), 0.0076)
        self.assertAlmostEqual(float(prices.max()), 1.7936)
        self.assertAlmostEqual(float(prices.mean()), 0.76620, places=5)

    def test_price_forecast_metrics_and_causality(self):
        audit = self.data.price.audit
        self.assertAlmostEqual(audit["seasonal_mae"], 0.04715, places=5)
        self.assertLess(audit["elastic_net_mae"], audit["seasonal_mae"])
        history = self.data.actual_prices.copy()
        original = feature_matrix(history, 100)
        history[100:] = 999.0
        np.testing.assert_array_equal(original, feature_matrix(history, 100))
        self.assertTrue(np.isnan(self.data.price.residuals[:7]).all())

    def test_prefix_uses_realized_slots_only(self):
        day, hour = 100, 12
        original = self.data.prefix(day, hour)
        actual = self.data.q3.actual[day, hour * 6:].copy()
        prices = self.data.actual_prices[day, hour * 6:].copy()
        try:
            self.data.q3.actual[day, hour * 6:] += 1e6
            self.data.actual_prices[day, hour * 6:] += 1e6
            np.testing.assert_array_equal(original, self.data.prefix(day, hour))
        finally:
            self.data.q3.actual[day, hour * 6:] = actual
            self.data.actual_prices[day, hour * 6:] = prices

    def test_joint_history_and_tail_reduction(self):
        config = Config(scenarios=10, bootstrap_repetitions=3, seconds=1, solver_threads=1)
        scenarios = construct(self.data, 60, 6, 144, config)
        self.assertEqual(scenarios.net.shape, scenarios.prices.shape)
        self.assertEqual(scenarios.distance.shape, (len(scenarios.weights),) * 2)
        np.testing.assert_allclose(np.diag(scenarios.distance), 0, atol=1e-12)
        self.assertAlmostEqual(float(scenarios.weights.sum()), 1)
        net = np.zeros((20, 4)); price = np.zeros((20, 4))
        net[-1] = 100; price[-1] = 10
        representatives, _, _, _, audit = _reduce(net, price, np.full(20, 0.05),
                                                   Config(scenarios=5, bootstrap_repetitions=1))
        self.assertIn(19, representatives)
        self.assertIn(19, audit["tail_indices_in_pool"])

    def test_physical_call_separation(self):
        policy = Policy(np.array([100.0, 100.0]), np.zeros(2), np.zeros(2))
        result = replay(policy, np.array([20.0, -10.0]), 6000.0)
        np.testing.assert_allclose(result["called"], [20, 0])
        np.testing.assert_allclose(result["unused"], [80, 100])
        np.testing.assert_allclose(result["spill"], [0, 10])
        np.testing.assert_allclose(result["unused"] + result["spill"], result["total_surplus"])

    def test_adjustment_settlement(self):
        grid = np.array([150.0, 50.0, 100.0])
        reference = np.full(3, 100.0)
        np.testing.assert_allclose(adjustment_quantity(grid, reference), [75.0, -25.0, 0.0])

    def test_dual_and_constraint_generation_agree(self):
        net = np.array([[1000., 1200, 800, 500], [1300., 1400, 900, 700]])
        prices = np.array([[.5, .6, .8, 1.], [.6, .8, 1., 1.2]])
        scenarios = Scenarios(net, prices, np.array([.6, .4]), np.array([[0., 1.], [1., 0.]]), .1, {})
        config = Config(scenarios=2, bootstrap_repetitions=1, seconds=5, gap=0,
                        solver_threads=1, allow_limited=False)
        generated = solve(scenarios, 1200.0, config)
        dual = solve_dual(scenarios, 1200.0, config)
        self.assertAlmostEqual(generated.audit["objective"], dual.audit["objective"], places=5)
        self.assertTrue(generated.audit["reliable"] and dual.audit["reliable"])

    def test_experiment_matrix_covers_required_ablations(self):
        variants = experiment_variants(Config(bootstrap_repetitions=1))
        required = {"independent_empirical", "joint_empirical", "joint_dro", "q43_unconditional",
                    "q43_conditional", "radius_0.75", "radius_1.25", "scenarios_10", "scenarios_30",
                    "net_weight_0.25", "net_weight_0.75", "window_84", "window_112",
                    "P0_0", "P1_06", "P2_0612", "P3_061218"}
        self.assertTrue(required.issubset(variants))

    def test_workbook_roundtrip_and_partial_guard(self):
        frame, _ = _cold_start(self.data, 0, 6000.0, "4-2")
        config = Config(bootstrap_repetitions=1)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            audit = export_workbook(frame, output, "4-2", config, smoke=True)
            self.assertTrue(audit["roundtrip_all_sheets"])
            workbook = load_workbook(output / "smoke_result4-2.xlsx", read_only=True, data_only=True)
            self.assertEqual(workbook.sheetnames, ["计划购电量", "充放电量", "紧急购电量"])
            workbook.close()
        with self.assertRaises(ValueError):
            validate_frame(frame, "4-2", config, full=True)

    def test_every_nonempty_source_line_is_mapped(self):
        with tempfile.TemporaryDirectory() as directory:
            audit = generate(Path(directory))
        self.assertTrue(audit["complete"])
        self.assertEqual(audit["mapped_lines"], audit["nonempty_lines"])


if __name__ == "__main__":
    unittest.main()
