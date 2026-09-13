"""Local audit experiments for 临时方案.md; not a competition submission solver.

Run: uv run --with scipy --with openpyxl python src/问题1/review_temporary_plan.py
Audit convention: input timestamps label the END of a ten-minute interval;
power is constant within that interval; charge/discharge are bus-side energies.
The original template has a conflicting interval labeling, so these numbers
are conditional review evidence, not certified result1.xlsx answers.
"""

import hashlib
import json
from pathlib import Path

import numpy as np
import scipy
from openpyxl import load_workbook
from scipy.optimize import Bounds, LinearConstraint, linprog, milp
from scipy.sparse import csc_matrix


ROOT = Path(__file__).resolve().parents[2]
BASE = ROOT / "docs/CUMCM2026Problems/C题/附件"
OUT = ROOT / "outputs/q1/raw"
DT = 1 / 6


def solve(prices, load, pv, initial=6000, eta=0.9, eta_d=None, lower=1200,
          upper=10800, power=5000, secondary=None, integer=False):
    n = len(prices)
    eta_d = eta if eta_d is None else eta_d
    size = 5 * n + 1
    cost = np.zeros(size)
    cost[:n] = prices
    eq = np.zeros((2 * n + 1, size))
    rhs = np.zeros(2 * n + 1)
    for t in range(n):
        eq[t, [t, n + t, 2 * n + t, 3 * n + t]] = [1, -1, 1, -1]
        rhs[t] = load[t] - pv[t]
        eq[n + t, [n + t, 2 * n + t, 4 * n + t, 4 * n + t + 1]] = [-eta, 1 / eta_d, -1, 1]
    eq[-1, [4 * n, 5 * n]] = [1, -1]
    bounds = ([(0, None)] * n + [(0, power * DT)] * (2 * n)
              + [(0, float(s)) for s in pv] + [(lower, upper)] * (n + 1))
    if initial is not None:
        bounds[4 * n] = (initial, initial)
    first = linprog(cost, A_eq=eq, b_eq=rhs, bounds=bounds, method="highs")
    if not first.success:
        raise RuntimeError(first.message)
    result = first
    if secondary is not None:
        objective = np.zeros(size)
        if secondary == "throughput":
            objective[n:3 * n] = 1
        elif secondary == "initial_min":
            objective[4 * n] = 1
        elif secondary == "initial_max":
            objective[4 * n] = -1
        result = linprog(objective, A_eq=np.vstack([eq, cost]),
                         b_eq=np.r_[rhs, first.fun], bounds=bounds, method="highs")
        if not result.success:
            raise RuntimeError(result.message)
    if integer:
        full_eq = np.pad(eq, ((0, 0), (0, n)))
        mutex = np.zeros((2 * n, size + n))
        for t in range(n):
            mutex[t, n + t] = 1
            mutex[t, size + t] = -power * DT
            mutex[n + t, 2 * n + t] = 1
            mutex[n + t, size + t] = power * DT
        low = [b[0] for b in bounds] + [0] * n
        high = [np.inf if b[1] is None else b[1] for b in bounds] + [1] * n
        result = milp(np.r_[cost, np.zeros(n)], integrality=np.r_[np.zeros(size), np.ones(n)],
                      bounds=Bounds(low, high), constraints=[
                          LinearConstraint(csc_matrix(full_eq), rhs, rhs),
                          LinearConstraint(csc_matrix(mutex), -np.inf,
                                           np.r_[np.zeros(n), np.full(n, power * DT)])],
                      options={"time_limit": 30, "mip_rel_gap": 1e-9})
        if not result.success:
            raise RuntimeError(result.message)
    x = result.x[:size]
    g, c, d, w, e = x[:n], x[n:2*n], x[2*n:3*n], x[3*n:4*n], x[4*n:]
    finite_upper = np.array([np.inf if b[1] is None else b[1] for b in bounds])
    dual = float(rhs @ first.eqlin.marginals + np.array([b[0] for b in bounds]) @ first.lower.marginals
                 + finite_upper[np.isfinite(finite_upper)] @ first.upper.marginals[np.isfinite(finite_upper)])
    summary = {
        "cost_yuan": float(cost @ x), "grid_kwh": float(g.sum()),
        "charge_kwh": float(c.sum()), "discharge_kwh": float(d.sum()),
        "curtail_kwh": float(w.sum()), "initial_kwh": float(e[0]), "final_kwh": float(e[-1]),
        "state_min": float(e.min()), "state_max": float(e.max()),
        "simultaneous_intervals": int(((c > 1e-7) & (d > 1e-7)).sum()),
        "balance_residual": float(np.max(np.abs(eq @ x - rhs))),
        "bound_violation": float(max(0, np.max(np.array([b[0] for b in bounds]) - x), np.max(x - finite_upper))),
        "cycle_residual": float(abs(d.sum() - eta*eta_d * c.sum())),
        "global_energy_residual": float(abs(g.sum() - (load.sum()-pv.sum()+w.sum()+(1-eta*eta_d)*c.sum()))),
        "primary_lp_dual_gap": abs(float(first.fun)-dual),
        "upper_state_raw_sum_including_fixed_initial": float(first.upper.marginals[4*n:].sum()),
        "upper_capacity_marginal": float(first.upper.marginals[4*n+(initial is not None):].sum()),
        "lower_capacity_marginal": float(first.lower.marginals[4*n+(initial is not None):].sum()),
        "initial_bound_marginal": float(first.lower.marginals[4*n]+first.upper.marginals[4*n]),
        "power_marginal_yuan_per_kw": float(DT * first.upper.marginals[n:3*n].sum()),
    }
    return summary, {"grid": g.tolist(), "charge": c.tolist(), "discharge": d.tolist(),
                     "curtail": w.tolist(), "state": e.tolist()}


def main():
    ws = load_workbook(BASE / "附件1.xlsx", read_only=True, data_only=True).active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    p, load, pv = (np.array([r[i] for r in rows], dtype=float) for i in (1, 2, 3))
    load, pv = load * DT, pv * DT
    template = load_workbook(BASE / "附件5/result1.xlsx", read_only=True)["计划购电量"]
    report = {"scipy_version": scipy.__version__, "assumptions": __doc__,
              "source_sha256": {str(f.relative_to(ROOT)): hashlib.sha256(f.read_bytes()).hexdigest()
                                 for f in [ROOT / "临时方案.md", BASE / "附件1.xlsx", BASE / "附件5/result1.xlsx"]},
              "input": {"intervals": len(rows), "price_min": float(p.min()), "price_max": float(p.max()),
                        "load_kwh": float(load.sum()), "pv_kwh": float(pv.sum()),
                        "pv_surplus_kwh": float(np.maximum(pv-load, 0).sum()),
                        "template_first": template["A2"].value, "template_last": template["A145"].value},
              "cases": {}, "schedules": {}}
    cases = {
        "fixed6000_eta90": {}, "free_initial_eta90": {"initial": None},
        "fixed6000_roundtrip90": {"eta": np.sqrt(0.9)},
        "roundtrip90_split_c90_d100": {"eta": 0.9, "eta_d": 1.0},
        "roundtrip90_split_c100_d90": {"eta": 1.0, "eta_d": 0.9},
        "fixed6000_secondary": {"secondary": "throughput"}, "fixed6000_milp": {"integer": True},
        "free_optimal_initial_min": {"initial": None, "secondary": "initial_min"},
        "free_optimal_initial_max": {"initial": None, "secondary": "initial_max"},
        "upper_plus_1": {"upper": 10801}, "upper_minus_1": {"upper": 10799},
        "nominal_plus_1_fixed6000": {"lower": 1200.1, "upper": 10800.9},
        "power_plus_1": {"power": 5001}, "power_minus_1": {"power": 4999},
        "wrong_5000kwh_per_slot": {"power": 30000},
        "wrong_allow_12000kwh": {"upper": 12000},
    }
    for name, options in cases.items():
        report["cases"][name], report["schedules"][name] = solve(p, load, pv, **options)
    for name in ["wrong_5000kwh_per_slot", "wrong_allow_12000kwh"]:
        s = report["schedules"][name]
        report["cases"][name]["against_original_limits"] = {
            "max_charge_kw": max(s["charge"]) / DT, "max_discharge_kw": max(s["discharge"]) / DT,
            "max_stored_kwh": max(s["state"])}
    report["cases"]["cyclic_timestamp_shift_10min"], _ = solve(np.roll(p,1), np.roll(load,1), np.roll(pv,1))
    report["baseline"] = {"without_storage_cost": float(p @ np.maximum(load-pv, 0)),
                          "without_storage_grid_kwh": float(np.maximum(load-pv, 0).sum())}
    # A synthetic one-slot falsification test, not a claim about attachment 4 prices.
    report["negative_price_counterexample"], _ = solve(np.array([-1.]), np.array([0.]), np.array([0.]),
                                                       initial=0, lower=0, upper=1, power=6)
    # A synthetic feasible but dissipative optimum: free PV can be curtailed instead.
    report["free_pv_cycle_example"] = {"pv_kwh": 0.19, "load_kwh": 0,
                                       "charge_kwh": 1, "discharge_kwh": 0.81,
                                       "grid_kwh": 0, "state_change_kwh": 0,
                                       "alternative": "charge=discharge=0, curtail=0.19, same zero cost"}
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / "temporary_plan_review_evidence.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({k:v for k,v in report.items() if k != "schedules"}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
