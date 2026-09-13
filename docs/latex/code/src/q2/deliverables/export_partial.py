"""只读导出已验收的前94日快照；不调用全年result2发布入口。"""

import datetime as dt
import hashlib
import html
import json
from pathlib import Path
import shutil
import sys
import tempfile
import zipfile

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from checkpoint import read_checkpoint, TABLES
from data import ROOT, sha256, write_csv, write_json
from dispatch import load_prepared
from export_dispatch import aggregate_storage, emergency_events


def main():
    run = ROOT / "outputs/q2/dispatch_runs/20260911_224501"
    source = run / "processed/main"
    output = run / "deliverables" / dt.datetime.now().strftime("partial_94days_%Y%m%d_%H%M%S")
    output.mkdir(parents=True)
    _, _, actual, dates, prices = load_prepared(run)
    # 活动原目录绝不调用恢复函数；只复制完整提交后在临时目录验收。
    with tempfile.TemporaryDirectory() as temp:
        copy = Path(temp)
        for _ in range(20):
            checkpoint = (source / "checkpoint.json").read_bytes()
            manifest = json.loads(checkpoint)
            tables = {name: (source / name).read_bytes() for name in TABLES}
            if (checkpoint == (source / "checkpoint.json").read_bytes()
                    and all(hashlib.sha256(raw).hexdigest() == manifest["sha256"][name]
                            for name, raw in tables.items())):
                break
        else:
            raise RuntimeError("没有读到稳定的已提交存档")
        assert manifest["days"] >= 94
        for name, raw in tables.items():
            (copy / name).write_bytes(raw)
        (copy / "checkpoint.json").write_bytes(checkpoint)
        status = json.loads((source / "run_status.json").read_text(encoding="utf-8"))
        write_json(copy / "run_status.json", status)
        saved_days = pd.read_csv(copy / "daily.csv", float_precision="round_trip")
        for folder, extension in (("daily_audit", ".json"), ("frozen_plans", ".csv")):
            (copy / folder).mkdir()
            for date in saved_days.date:
                shutil.copyfile(source / folder / (date + extension), copy / folder / (date + extension))
        frame, daily, calibration = read_checkpoint(
            copy, actual, dates, prices, float(saved_days.initial_soc.iloc[0]),
            status["start_index"], status["end_index"], status["settings"])
        frame = frame.iloc[:94 * 144].copy()
        daily = pd.DataFrame(daily).iloc[:94].copy()
        selected_dates = set(daily.date)
        calibration = pd.DataFrame(calibration)
        calibration = calibration[calibration.origin.isin(selected_dates)].copy()
        audits = []
        for date in daily.date:
            audit = json.loads((copy / f"daily_audit/{date}.json").read_text(encoding="utf-8"))
            solver = audit[str(audit["selected_K"])]["selected_solver"]
            audits.append({"date": date, "requested_gap": solver.get("requested_gap", .01),
                           "achieved_gap": solver["mip_gap"], "selected_reliable": solver["reliable"],
                           "computational_limited": audit["computational_limited"],
                           "audit_sha256": sha256(copy / f"daily_audit/{date}.json")})
        audit_table = pd.DataFrame(audits)
    assert len(frame) == 13536 and len(daily) == 94
    assert daily.date.iloc[0] == "2025-02-01" and daily.date.iloc[-1] == "2025-05-05"
    assert audit_table.selected_reliable.all()
    assert (audit_table.achieved_gap <= audit_table.requested_gap + 1e-8).all()
    tables = {"dispatch": frame, "daily": daily, "calibration": calibration,
              "storage_blocks": aggregate_storage(frame), "emergency_events": emergency_events(frame),
              "solver_acceptance": audit_table}
    for name, table in tables.items():
        write_csv(output / f"{name}.csv", table)
    explanation = [
        ("文件性质", "第二问已完成94天的阶段结果；不是全年最终result2"),
        ("日期范围", "2025-02-01 至 2025-05-05"),
        ("已保存模拟日", 94), ("10分钟调度记录", 13536),
        ("精度记录", "前85天采用1%；后9天采用2%；逐日门槛与实际gap见求解验收"),
        ("验收范围", "原始输入、电价、连续日期、跨日SOC、原反馈物理与费用逐值核验通过"),
        ("候选限制", "选中策略达标；部分未选候选曾超时，computational_limited标识如实保留"),
        ("数值性质", "冻结仿真输出，数值逐值导入CSV；不是可编辑的Excel计算模型，无Excel公式"),
        ("单位", "电量和SOC为kWh，负荷与光伏功率为kW，费用为元；单价为元/kWh"),
        ("时间定义", "timestamp为10分钟区间右端点；24:00归属前一源日期"),
        ("剩余工作", "其余240个模拟日及尚未完成实验不在本文件内；正式计算继续"),
        ("来源", str(source.relative_to(ROOT))),
        ("快照生成时间", dt.datetime.now().astimezone().isoformat()),
    ]
    labels = {"date": "日期", "K": "预测跨度（天）", "M": "历史块数", "S": "场景数",
              "initial_soc": "日初电量（kWh）", "final_soc": "日末电量（kWh）", "cost": "实际日费用（元）",
              "emergency_kwh": "紧急购电（kWh）", "computational_limited": "存在候选计算受限",
              "interval": "时段", "timestamp": "区间右端点", "grid": "计划购电（kWh）",
              "charge_cap": "充电响应上限（kWh）", "discharge_cap": "放电响应上限（kWh）",
              "net_kwh": "净负荷（kWh）", "price": "电价（元每kWh）", "charge": "实际充电（kWh）",
              "discharge": "实际放电（kWh）", "emergency": "紧急购电（kWh）", "unused": "未利用电量（kWh）",
              "soc": "时段末电量（kWh）", "planned_cost": "计划购电费用（元）", "emergency_cost": "紧急购电费用（元）",
              "predicted_load_kw": "预测负荷（kW）", "predicted_pv_kw": "预测光伏（kW）"}
    workbook = Workbook()
    workbook.remove(workbook.active)
    contents = {"说明": pd.DataFrame(explanation, columns=["项目", "说明"]), "每日汇总": daily,
                "10分钟调度明细": frame, "4小时充放电": tables["storage_blocks"],
                "紧急购电区间": tables["emergency_events"], "求解验收": audit_table}
    for title, table in contents.items():
        sheet = workbook.create_sheet(title)
        sheet.append([labels.get(key, key) for key in table.columns])
        for row in table.itertuples(index=False, name=None):
            sheet.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in row])
        for row in sheet:
            for cell in row:
                cell.font = Font(name="Arial", size=10)
                if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0.000000"
                if isinstance(cell.value, dt.datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="163D52")
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        for column in range(1, len(table.columns) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 25
        sheet.freeze_panes = "C2" if title == "10分钟调度明细" else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
    workbook["说明"].column_dimensions["B"].width = 105
    for row in workbook["说明"].iter_rows(min_row=2):
        row[1].alignment = Alignment(wrap_text=True, vertical="center")
        workbook["说明"].row_dimensions[row[0].row].height = 34
    filename = output / "第二问_94天阶段结果_20250201-20250505.xlsx"
    workbook.save(filename); workbook.close()
    check = load_workbook(filename, read_only=True, data_only=True)
    for title, original in (("每日汇总", daily), ("10分钟调度明细", frame)):
        rows = list(check[title].values)
        assert len(rows) == len(original) + 1
        restored = pd.DataFrame(rows[1:], columns=original.columns)
        numeric = original.select_dtypes(include="number").columns
        np.testing.assert_allclose(restored[numeric].to_numpy(dtype=float), original[numeric].to_numpy(dtype=float),
                                   atol=1e-8, rtol=1e-12)
    assert not any(cell.data_type == "e" for sheet in check for row in sheet for cell in row)
    check.close()
    metadata = {"partial": True, "days": 94, "records": len(frame), "start": daily.date.iloc[0], "end": daily.date.iloc[-1],
                "source_checkpoint": manifest, "original_data_physics_cost_and_soc_validated": True,
                "excel_roundtrip_passed": True, "workbook_sha256": sha256(filename),
                "csv_sha256": {name: sha256(output / f"{name}.csv") for name in tables}}
    write_json(output / "manifest.json", metadata)
    preview = "<!doctype html><meta charset='utf-8'><title>第二问94天阶段结果</title><style>body{font:16px Arial;max-width:1100px;margin:36px auto;color:#163d52}td,th{padding:8px;border-bottom:1px solid #ddd}table{border-collapse:collapse}a{color:#066e85}</style>"
    preview += f"<h1>第二问：已完成94天的阶段结果</h1><p>2025-02-01 至 2025-05-05 · 13,536条10分钟记录 · 存档核验与Excel逐值回读通过</p><p><a href='{html.escape(filename.name)}'>下载94天结果 Excel</a></p><p>这是阶段快照，全年尚未算完。前85天采用1%，后9天采用2%，历史候选受限记录保留。</p>"
    preview += "<h2>前5日结果预览</h2>" + daily.head().rename(columns=labels).to_html(index=False)
    (output / "index.html").write_text(preview, encoding="utf-8")
    with zipfile.ZipFile(output / "第二问_94天结果及CSV.zip", "w", zipfile.ZIP_DEFLATED) as bundle:
        for path in output.iterdir():
            if path.suffix != ".zip":
                bundle.write(path, path.name)
    print(json.dumps({"directory": str(output), "xlsx": str(filename), **metadata}, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
