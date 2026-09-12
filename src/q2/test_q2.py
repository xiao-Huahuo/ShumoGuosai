"""真实附件验收 + 手算/合成数据检验；合成结果不作为实际预测交付。"""

import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

import data
from data import (CALENDAR, FORECAST_START, ROOT, SERIES, SHEETS, calendar_for_day,
                  endpoint_minutes, history_before, jump_context, read_attachment)
from backtest import dispatch_metrics, forecast_metrics, rolling_folds, rolling_predict
from diagnostics import (correlation, describe, lag_analysis, seasonal_transforms, spectral_analysis,
                         stability_analysis, stationarity_analysis)
from report import candidate_models


class Q2Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = ROOT / "docs/CUMCM2026Problems/C题/附件/附件2.xlsx"
        cls.frame, cls.qc = read_attachment(cls.path, Path(cls.tmp.name) / "quality.csv")
        workbook = load_workbook(cls.path, read_only=True, data_only=True)
        cls.rows = {s.title: list(s.iter_rows(values_only=True)) for s in workbook}
        workbook.close()

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_original_values_and_quality(self):
        self.assertTrue(self.qc.status.eq("pass").all())
        self.assertEqual(len(self.frame), 52560)
        for variable, name in SHEETS.items():
            original = np.array([row[1:] for row in self.rows[name][1:]], dtype=float).ravel()
            np.testing.assert_array_equal(original, self.frame[variable])
        np.testing.assert_array_equal(self.frame.net_load_kw, self.frame.load_kw - self.frame.pv_kw)
        self.assertGreater(int((self.frame.pv_kw == 0).sum()), 0)
        self.assertGreater(int((self.frame.net_load_kw < 0).sum()), 0)

    def test_right_endpoint_and_calendar(self):
        first, last = self.frame.iloc[0], self.frame.iloc[-1]
        self.assertEqual(first.timestamp, pd.Timestamp("2025-01-01 00:10"))
        self.assertEqual(last.timestamp, pd.Timestamp("2026-01-01"))
        self.assertEqual(last.date, pd.Timestamp("2025-12-31"))
        self.assertEqual(last.interval, "23:50-24:00")
        self.assertEqual(self.frame.iloc[143].weekday, 3)
        self.assertTrue(self.frame.timestamp.diff().dropna().eq(pd.Timedelta(minutes=10)).all())
        for value in ("0:00+1", "24:00", 1.0):
            self.assertEqual(endpoint_minutes(value), 1440)
        self.assertEqual(endpoint_minutes(1/144), 10)
        for value in ("0:00", "24:10", "0:11", float("nan"), True):
            with self.assertRaises((ValueError, TypeError)):
                endpoint_minutes(value)

    def assert_bad_workbook(self, mutate):
        rows = {name: [list(row) for row in values] for name, values in self.rows.items()}
        mutate(rows)
        class Sheet:
            def __init__(self, values):
                self.values = values
            def iter_rows(self, **kwargs):
                return iter(self.values)
        class Book:
            sheetnames = list(rows)
            def __getitem__(self, key):
                return Sheet(rows[key])
            def close(self):
                pass
        with patch.object(data, "load_workbook", return_value=Book()):
            with self.assertRaises(ValueError):
                read_attachment(self.path, Path(self.tmp.name) / "bad.csv")
        self.assertIn("fail", pd.read_csv(Path(self.tmp.name) / "bad.csv").status.values)

    def test_missing_and_nonfinite_rejected(self):
        for invalid in (None, float("nan"), float("inf"), "unreadable", True):
            with self.subTest(invalid=invalid):
                self.assert_bad_workbook(lambda rows: rows["小区负载"][10].__setitem__(8, invalid))

    def test_negative_physical_values_rejected(self):
        for sheet in SHEETS.values():
            self.assert_bad_workbook(lambda rows: rows[sheet][10].__setitem__(8, -1))

    def test_duplicate_and_missing_dates_rejected(self):
        self.assert_bad_workbook(lambda rows: rows["小区负载"][3].__setitem__(0, rows["小区负载"][2][0]))
        self.assert_bad_workbook(lambda rows: rows["小区负载"].pop())

    def test_duplicate_missing_and_wrong_endpoints_rejected(self):
        self.assert_bad_workbook(lambda rows: rows["光伏发电实际功率"][0].__setitem__(5, rows["光伏发电实际功率"][0][4]))
        self.assert_bad_workbook(lambda rows: rows["光伏发电实际功率"][0].__setitem__(144, "0:00"))

    def test_missing_sheet_rejected(self):
        self.assert_bad_workbook(lambda rows: rows.pop("光伏发电实际功率"))

    def test_jump_review_preserves_spikes_and_uses_past_peers(self):
        part = self.frame.iloc[:3000].copy()
        part.loc[2016, "load_kw"] = 1000000.
        before = part.copy(deep=True)
        context = jump_context(part)
        pd.testing.assert_frame_equal(part, before)
        row = context.loc[(context.series == "load_kw") & (context.timestamp == part.loc[2016, "timestamp"])].iloc[0]
        self.assertEqual(row.value_kw, 1000000.)
        self.assertEqual(row.previous_week_kw, part.loc[1008, "load_kw"])
        self.assertEqual(row.past_same_weekday_count, 2)
        self.assertEqual(row.past_same_weekday_max, max(part.loc[0, "load_kw"], part.loc[1008, "load_kw"]))
        self.assertEqual(row.action, "retain_review_only")

    def test_strict_midnight_and_334_folds(self):
        folds = rolling_folds(self.frame)
        self.assertEqual(len(folds), 334)
        self.assertEqual(folds.iloc[0].history_points, 4463)
        self.assertEqual(folds.iloc[0].history_last_timestamp, pd.Timestamp("2025-01-31 23:50"))
        self.assertTrue((folds.history_last_timestamp < folds.origin).all())
        self.assertTrue(folds.horizon_points.eq(144).all())
        self.assertEqual(folds.iloc[-1].target_last, pd.Timestamp("2026-01-01"))
        with self.assertRaises(ValueError):
            history_before(self.frame, "2025-02-01 00:10")
        with self.assertRaises(ValueError):
            rolling_folds(self.frame.iloc[:-1])

    @staticmethod
    def mean_factory(variable):
        def predict(history, future):
            assert set(history) == {"timestamp", "value_kw"}
            assert tuple(future.columns) == CALENDAR
            assert history.timestamp.max() < future.date.iloc[0]
            return np.full(144, history.value_kw.mean())
        return predict

    def test_future_and_midnight_perturbation_cannot_change_prediction(self):
        a = rolling_predict(self.frame, self.mean_factory, end=FORECAST_START)
        altered = self.frame.copy()
        mask = altered.timestamp >= FORECAST_START
        altered.loc[mask, "load_kw"] += 10000
        altered.loc[mask, "pv_kw"] += 20000
        altered["net_load_kw"] = altered.load_kw - altered.pv_kw
        b = rolling_predict(altered, self.mean_factory, end=FORECAST_START)
        for v in SERIES:
            np.testing.assert_array_equal(a[f"predicted_{v}"], b[f"predicted_{v}"])
        self.assertFalse(a.actual_load_kw.equals(b.actual_load_kw))
        # 反向检查：可见历史改变，预测应改变；避免恒定假预测使防泄漏测试虚假通过。
        altered.loc[altered.timestamp < FORECAST_START, "load_kw"] += 500
        c = rolling_predict(altered, self.mean_factory, end=FORECAST_START)
        np.testing.assert_allclose(c.predicted_load_kw - a.predicted_load_kw, 500)

    def test_new_model_per_variable_per_day(self):
        calls = []
        def factory(variable):
            calls.append(variable)
            return self.mean_factory(variable)
        result = rolling_predict(self.frame, factory, end="2025-02-02")
        self.assertEqual(calls, ["load_kw", "pv_kw"] * 2)
        self.assertEqual(len(result), 288)
        np.testing.assert_array_equal(result.predicted_net_load_kw, result.predicted_load_kw - result.predicted_pv_kw)

    def test_invalid_predictions_rejected(self):
        for values in (np.ones(143), np.full(144, np.nan), np.full(144, -1)):
            with self.assertRaises(ValueError):
                rolling_predict(self.frame, lambda _: lambda past, future: values, end=FORECAST_START)

    def test_forecast_metrics_hand_calculation(self):
        frame = calendar_for_day("2025-02-01")
        for v, actual, predicted in (("load_kw", 10, np.tile([12, 8], 72)),
                                     ("pv_kw", 4, np.tile([5, 2], 72))):
            frame[f"actual_{v}"] = actual
            frame[f"predicted_{v}"] = predicted
        frame["actual_net_load_kw"] = 6
        frame["predicted_net_load_kw"] = frame.predicted_load_kw - frame.predicted_pv_kw
        all_metrics = forecast_metrics(frame).query("period == 'all'").set_index("series")
        self.assertEqual(all_metrics.loc["load_kw", "mae_kw"], 2)
        self.assertEqual(all_metrics.loc["pv_kw", "mae_kw"], 1.5)
        self.assertAlmostEqual(all_metrics.loc["pv_kw", "rmse_kw"], np.sqrt(2.5))
        self.assertEqual(all_metrics.loc["net_load_kw", "mae_kw"], .5)
        self.assertAlmostEqual(all_metrics.loc["net_load_kw", "rmse_kw"], np.sqrt(.5))
        frame.loc[0, "predicted_net_load_kw"] = 999
        with self.assertRaises(ValueError):
            forecast_metrics(frame)

    def test_dispatch_metrics_plan_billing_and_fivefold_emergency(self):
        frame = calendar_for_day("2025-02-01")
        frame["planned_grid_kwh"], frame["price_yuan_per_kwh"] = 2., .5
        frame["actual_load_kw"], frame["actual_supply_before_emergency_kw"] = 12., 12.
        frame.loc[[0, 1, 3], "actual_supply_before_emergency_kw"] = [0., 6., 0.]
        m = dispatch_metrics(frame)
        self.assertEqual(m["emergency_kwh"], 5)
        self.assertEqual(m["emergency_intervals"], 3)
        self.assertEqual(m["emergency_events"], 2)
        self.assertEqual(m["days_with_emergency"], 1)
        self.assertEqual(m["planned_cost_yuan"], 144)
        self.assertEqual(m["emergency_cost_yuan"], 12.5)
        self.assertEqual(m["total_cost_yuan"], 156.5)
        self.assertEqual(m["cost_period_label"], "evaluated_period_only")
        frame.actual_supply_before_emergency_kw = 999
        self.assertEqual(dispatch_metrics(frame)["total_cost_yuan"], 144)  # 多余供给不能退款。
        frame.loc[0, "actual_supply_before_emergency_kw"] = -6.
        self.assertEqual(dispatch_metrics(frame)["emergency_kwh"], 3.)  # 充电可能令紧急补足前净供给为负。
        with self.assertRaises(ValueError):
            dispatch_metrics(frame.iloc[:-1])

    def test_cross_midnight_emergency_is_one_contiguous_event(self):
        frame = pd.concat([calendar_for_day(d) for d in ("2025-02-01", "2025-02-02")], ignore_index=True)
        frame["planned_grid_kwh"], frame["price_yuan_per_kwh"] = 0., 1.
        frame["actual_load_kw"], frame["actual_supply_before_emergency_kw"] = 0., 0.
        frame.loc[[143, 144], "actual_load_kw"] = 6.
        m = dispatch_metrics(frame)
        self.assertEqual(m["emergency_events"], 1)
        self.assertEqual(m["days_with_emergency"], 2)

    def test_transforms_match_four_term_identity(self):
        x = np.arange(3000, dtype=float) ** 2
        transforms = seasonal_transforms(x)
        self.assertEqual(transforms["diff144"].notna().sum(), 3000-144)
        np.testing.assert_array_equal(transforms["diff144_then1008"].dropna(),
                                      x[1152:] - x[1008:-144] - x[144:-1008] + x[:-1152])

    def test_spectrum_known_daily_weekly_peaks_and_parseval(self):
        frame = pd.concat([calendar_for_day(d) for d in pd.date_range("2025-01-01", periods=28)], ignore_index=True)
        t = np.arange(len(frame))/144
        frame["load_kw"] = 10 + 2*np.cos(2*np.pi*t) + np.cos(2*np.pi*t/7)
        frame["pv_kw"] = 5 + 2*np.cos(2*np.pi*t)
        _, targets = spectral_analysis(frame, "test")
        load = targets.query("series == 'load_kw' and form == 'raw'")
        self.assertTrue(load.local_peak.all())
        np.testing.assert_allclose(load.integrated_psd_kw2, 2.5, atol=1e-12)
        np.testing.assert_allclose(targets.integrated_psd_kw2, targets.population_variance_kw2, atol=1e-12)

    def test_acf_and_pacf_computed_through_weekly_lags(self):
        part = history_before(self.frame, FORECAST_START)
        keys, curves = lag_analysis(part, "initial_history")
        self.assertEqual(curves.groupby("series").lag.max().tolist(), [2016, 2016])
        x = part.load_kw.to_numpy()
        centered = x-x.mean()
        actual = curves.query("series == 'load_kw' and lag == 1008").acf.iloc[0]
        self.assertAlmostEqual(actual, np.dot(centered[1008:], centered[:-1008])/np.dot(centered, centered))
        self.assertTrue(np.isfinite(curves.pacf).all())
        r = keys.query("series == 'load_kw' and form == 'raw' and lag == 144").pearson_r.iloc[0]
        self.assertAlmostEqual(r, np.corrcoef(x[144:], x[:-144])[0, 1])

    def test_monthly_and_weekday_profiles_and_stability(self):
        profiles, daily, summary, zeros = describe(self.frame)
        self.assertEqual(len(profiles.query("grouping == 'month'")), 12*144*3)
        self.assertEqual(len(profiles.query("grouping == 'weekday'")), 7*144*3)
        self.assertEqual(len(daily), 365*3)
        january = self.frame.loc[self.frame.month == 1]
        stability = stability_analysis(self.frame, "full_year")
        row = stability.query("month == 1 and series == 'load_kw'").iloc[0]
        self.assertEqual(row.weekly_pairs, 31*144-1008)
        self.assertAlmostEqual(row.weekly_lag_r, correlation(january.load_kw.to_numpy()[1008:], january.load_kw.to_numpy()[:-1008]))
        self.assertEqual(zeros.zero_points.sum(), int((self.frame.pv_kw == 0).sum()))

    def test_constant_series_not_reported_as_significant(self):
        part = self.frame.iloc[:2000].copy()
        part.load_kw, part.pv_kw = 1., 0.
        tests, _, _ = stationarity_analysis(part, "test")
        self.assertTrue(tests.status.eq("undefined_constant_or_short").all())
        self.assertTrue(np.isnan(correlation([1, 1], [2, 2])))

    def test_candidate_guard_rejects_full_year_information(self):
        with self.assertRaises(ValueError):
            candidate_models({"lag_correlations": pd.DataFrame({"scope": ["full_year"]}),
                              "spectral_targets": pd.DataFrame({"scope": ["full_year"]})})

    def test_small_weekly_peak_alone_does_not_prioritize_multiseason_model(self):
        history = history_before(self.frame, FORECAST_START)
        lags, _ = lag_analysis(history, "initial_history")
        _, targets = spectral_analysis(history, "initial_history")
        candidates = candidate_models({"lag_correlations": lags, "spectral_targets": targets})
        pv = candidates.query("series == 'pv_kw' and model == 'DHR_AR'").iloc[0]
        self.assertEqual(pv.status, "条件保留，不预设周效应")
        self.assertTrue(candidates.fitted.eq(False).all())

    @unittest.skipIf(os.name == "nt", "旧2.1发布器依赖目录符号链接；新最终模型使用独立运行目录")
    def test_failed_publish_keeps_current_generation(self):
        import run
        root = Path(self.tmp.name) / "atomic"
        previous, following = root / "outputs/.q2-runs/previous", root / "outputs/.q2-runs/following"
        for p in (previous, following):
            p.mkdir(parents=True, exist_ok=True)
        with patch.object(run, "ROOT", root):
            run.publish(previous)
            with patch.object(run.os, "replace", side_effect=OSError("injected publish failure")):
                with self.assertRaises(OSError):
                    run.publish(following)
        self.assertEqual((root / "outputs/q2_current").resolve(), previous.resolve())
        with patch.object(run, "ROOT", root):
            run.publish(following)
        self.assertEqual((root / "outputs/q2_current").resolve(), following.resolve())

    def test_generated_artifacts_roundtrip(self):
        generation = os.environ.get("Q2_GENERATION")
        if not generation:
            self.skipTest("仅完整流水线设置Q2_GENERATION后执行产物回读；其余测试独立运行")
        path = Path(generation)
        saved = pd.read_csv(path / "inputs/processed/timeseries.csv", float_precision="round_trip")
        for variable in SERIES:
            np.testing.assert_array_equal(saved[variable], self.frame[variable])
        keys = pd.read_csv(path / "processed/full_year_lag_correlations.csv")
        self.assertEqual(len(keys), 24)
        self.assertEqual(len(pd.read_csv(path / "processed/rolling_folds.csv")), 334)
        for scope in ("full_year", "initial_history"):
            tests = pd.read_csv(path / f"processed/{scope}_stationarity_tests.csv")
            self.assertEqual(len(tests), 32)
            self.assertTrue(tests.status.eq("ok").all())
        self.assertEqual(len(list((path / "processed/figures").glob("*.png"))), 11)
        self.assertEqual(len(list((path / "processed/figures").glob("*.svg"))), 11)
        trace = pd.read_csv(path / "processed/source_traceability.csv")
        source = (ROOT / "docs/2/dierwen1.md").read_text(encoding="utf-8").splitlines()
        self.assertEqual(set(trace.source_line), {i for i, line in enumerate(source, 1) if line.strip()})


if __name__ == "__main__":
    unittest.main(verbosity=2)
