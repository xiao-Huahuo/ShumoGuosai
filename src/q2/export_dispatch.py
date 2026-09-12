"""方案2.9.3：真实回放CSV、连续紧急区间及result2原模板回填/逐值回读。"""

from copy import copy
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment

from data import calendar_for_day, write_csv, write_json
from policy import CAP, E_MIN, E_MAX, ETA_D, PHYSICAL_TOL, validate_replay
from protocol import PROTOCOL


def validate_dispatch(frame, *, full_period=False, rigid=False):
    if frame.empty:
        raise ValueError("真实回放为空")
    required = ["grid", "net_kwh", "initial_soc", "charge", "discharge", "emergency", "unused", "soc", "price", "planned_cost", "emergency_cost"]
    if not np.isfinite(frame[required].to_numpy(dtype=float)).all() or (frame.grid < 0).any() or (frame.price <= 0).any():
        raise ValueError("回放包含非有限数值、负购电或非正电价")
    dates = pd.to_datetime(frame.date)
    times = pd.to_datetime(frame.timestamp)
    if times.duplicated().any() or not times.diff().dropna().eq(pd.Timedelta(minutes=10)).all():
        raise ValueError("真实回放时间必须连续唯一")
    if not dates.eq((times-pd.Timedelta(minutes=10)).dt.normalize()).all():
        raise ValueError("源日期与右端点不符")
    previous = None
    for date, day in frame.groupby(dates, sort=False):
        if len(day) != 144 or not np.array_equal(pd.to_datetime(day.timestamp).to_numpy(), calendar_for_day(date).timestamp.to_numpy()):
            raise ValueError("日期不是完整144时段")
        if previous is not None and abs(float(day.initial_soc.iloc[0])-previous) > PHYSICAL_TOL:
            raise ValueError("跨日SOC不连续")
        if not np.allclose(day.initial_soc, day.initial_soc.iloc[0], atol=PHYSICAL_TOL, rtol=0):
            raise ValueError("同一天的0:00初态不一致")
        reserve = day.soc_reserve_kwh.to_numpy() if "soc_reserve_kwh" in day else np.zeros(len(day))
        validate_replay(day.grid.to_numpy(), day.net_kwh.to_numpy(), float(day.initial_soc.iloc[0]),
                        {k: day[k].to_numpy() for k in ("charge", "discharge", "emergency", "unused", "soc")},
                        rigid=rigid, reserve=reserve)
        np.testing.assert_allclose(day.planned_cost, day.price*day.grid, atol=PHYSICAL_TOL, rtol=0)
        np.testing.assert_allclose(day.emergency_cost, 5*day.price*day.emergency, atol=PHYSICAL_TOL, rtol=0)
        previous = float(day.soc.iloc[-1])
    if full_period and (dates.min() != pd.Timestamp("2025-02-01") or dates.max() != pd.Timestamp("2025-12-31") or len(frame) != 334*144):
        raise ValueError("result2必须完整覆盖2025年2—12月，不能发布部分日期")


def emergency_events(frame):
    """q>0严格定义；连续区间跨午夜继续合并，电量与费用逐段累加。"""
    columns = ["date", "start", "end", "interval", "emergency_kwh", "cost_yuan", "intervals"]
    rows, current = [], None
    for row in frame.itertuples():
        end = pd.Timestamp(row.timestamp)
        start = end-pd.Timedelta(minutes=10)
        if row.emergency > 0:
            if current is None or start != current["end"]:
                if current is not None:
                    rows.append(current)
                current = {"date": start.normalize(), "start": start, "end": end,
                           "emergency_kwh": 0.0, "cost_yuan": 0.0, "intervals": 0}
            current["end"] = end
            current["emergency_kwh"] += row.emergency
            current["cost_yuan"] += row.emergency_cost
            current["intervals"] += 1
        elif current is not None:
            rows.append(current); current = None
    if current is not None:
        rows.append(current)
    for row in rows:
        days = (row["end"].normalize()-row["start"].normalize()).days
        suffix = "" if days == 0 else f"+{days}日"
        row["interval"] = f'{row["start"]:%H:%M}-{row["end"]:%H:%M}{suffix}'
    return pd.DataFrame(rows, columns=columns)


def aggregate_storage(frame):
    rows = []
    for date, day in frame.groupby("date", sort=False):
        for block in range(6):
            part = day.iloc[24*block:24*(block+1)]
            rows.append({"date": str(pd.Timestamp(date).date()), "interval": f"{block*4:02d}:00-{(block+1)*4:02d}:00",
                         "charge_kwh": float(part.charge.sum()), "discharge_kwh": float(part.discharge.sum()),
                         "soc_at_0000": float(day.initial_soc.iloc[0]), "soc_at_2400": float(day.soc.iloc[-1])})
    return pd.DataFrame(rows)


def summary(frame):
    events = emergency_events(frame)
    return {"start": str(pd.Timestamp(frame.date.iloc[0]).date()), "end": str(pd.Timestamp(frame.date.iloc[-1]).date()),
            "days": int(frame.date.nunique()), "period": "evaluation_period_not_calendar_year",
            "planned_cost_yuan": float(frame.planned_cost.sum()), "emergency_cost_yuan": float(frame.emergency_cost.sum()),
            "total_cost_yuan": float(frame.planned_cost.sum()+frame.emergency_cost.sum()),
            "planned_grid_kwh": float(frame.grid.sum()), "emergency_kwh": float(frame.emergency.sum()),
            "emergency_intervals": int((frame.emergency > 0).sum()), "emergency_events": len(events),
            "intervals_above_numerical_tolerance": int((frame.emergency > PHYSICAL_TOL).sum()),
            "unused_kwh": float(frame.unused.sum()), "charge_kwh": float(frame.charge.sum()),
            "discharge_kwh": float(frame.discharge.sum()), "final_soc": float(frame.soc.iloc[-1]),
            "charge_with_emergency_intervals": int(((frame.charge > 0)&(frame.emergency > 0)).sum())}


def write_tables(frame, output):
    write_csv(output/"emergency_events.csv", emergency_events(frame))
    storage = aggregate_storage(frame)
    write_csv(output/"storage_blocks.csv", storage)
    specified = frame[pd.to_datetime(frame.date).dt.strftime("%Y-%m-%d").isin(PROTOCOL["specified_dates"])]
    rows = []
    for date, day in specified.groupby("date"):
        for slot in (61, 73, 85, 97, 109, 121):
            point = day[day.slot == slot].iloc[0]
            rows.append({"date": str(pd.Timestamp(date).date()), "interval": point.interval,
                         "planned_kwh": float(point.grid), "daily_planned_kwh": float(day.grid.sum()),
                         "daily_planned_cost_yuan": float(day.planned_cost.sum()),
                         "daily_actual_cost_yuan": float(day.planned_cost.sum()+day.emergency_cost.sum())})
    write_csv(output/"paper_table1.csv", pd.DataFrame(rows, columns=["date", "interval", "planned_kwh", "daily_planned_kwh", "daily_planned_cost_yuan", "daily_actual_cost_yuan"]))
    write_csv(output/"paper_table2.csv", storage[storage.date.isin(PROTOCOL["specified_dates"])])
    # 指定日期表按日截取区间；全局事件表另外保留跨日连续事件定义。
    events = [table for _, day in specified.groupby("date")
              if not (table := emergency_events(day)).empty]
    write_csv(output/"paper_table3.csv", pd.concat(events, ignore_index=True) if events else emergency_events(specified))
    write_json(output/"summary.json", summary(frame))


def final_validation(frame):
    previous = np.r_[frame.initial_soc.iloc[0], frame.soc.to_numpy()[:-1]]
    balance = (frame.grid+frame.discharge+frame.emergency
               -frame.net_kwh-frame.charge-frame.unused)
    soc_violation = (frame.soc < E_MIN-PHYSICAL_TOL) | (frame.soc > E_MAX+PHYSICAL_TOL)
    simultaneous = (frame.charge > PHYSICAL_TOL) & (frame.discharge > PHYSICAL_TOL)
    power = (frame.charge > CAP+PHYSICAL_TOL) | (frame.discharge > CAP+PHYSICAL_TOL)
    emergency_charge = (frame.emergency > PHYSICAL_TOL) & (frame.charge > PHYSICAL_TOL)
    reserve = frame.soc_reserve_kwh.to_numpy() if "soc_reserve_kwh" in frame else np.zeros(len(frame))
    reserve_discharge = frame.discharge.to_numpy() > (
        ETA_D*np.maximum(previous-E_MIN-reserve, 0)+PHYSICAL_TOL)
    result = {
        "soc_min": float(frame.soc.min()), "soc_max": float(frame.soc.max()),
        "soc_violation_count": int(soc_violation.sum()),
        "simultaneous_charge_discharge_count": int(simultaneous.sum()),
        "power_violation_count": int(power.sum()),
        "max_energy_balance_residual": float(np.abs(balance).max()),
        "emergency_and_charge_count": int(emergency_charge.sum()),
        "reserve_discharge_violation_count": int(reserve_discharge.sum()),
        "intervals_checked": len(frame)}
    result["overall_pass"] = bool(
        result["soc_violation_count"] == 0
        and result["simultaneous_charge_discharge_count"] == 0
        and result["power_violation_count"] == 0
        and result["max_energy_balance_residual"] <= PHYSICAL_TOL
        and result["emergency_and_charge_count"] == 0
        and result["reserve_discharge_violation_count"] == 0)
    return result


def final_calibration(frame, actual, calibration):
    day_indices = pd.DatetimeIndex(pd.to_datetime(frame.date).unique()).dayofyear.to_numpy()-1
    realized = np.asarray(actual)[day_indices]
    predicted_load = frame.predicted_load_kw.to_numpy()
    predicted_pv = frame.predicted_pv_kw.to_numpy()
    actual_load, actual_pv = realized[..., 0].ravel(), realized[..., 1].ravel()
    series = {
        "load": (actual_load, predicted_load),
        "pv": (actual_pv, predicted_pv),
        "net_load": (actual_load-actual_pv, predicted_load-predicted_pv)}
    rows, payload = [], {}
    for name, (observed, predicted) in series.items():
        error = observed-predicted
        values = {"mae_kw": float(np.abs(error).mean()),
                  "rmse_kw": float(np.sqrt(np.mean(error**2))), "n": len(error)}
        payload[name] = values
        rows.extend([
            {"metric": f"{name}_MAE", "value": values["mae_kw"], "unit": "kW", "n": len(error)},
            {"metric": f"{name}_RMSE", "value": values["rmse_kw"], "unit": "kW", "n": len(error)}])
    calibration = calibration[calibration.horizon == 0]
    coverage = {}
    for nominal, label in ((.8, "PICP80"), (.9, "PICP90")):
        selected = calibration[np.isclose(calibration.nominal, nominal)]
        covered, n = int(selected.covered_intervals.sum()), int(selected.n.sum())
        value = covered/n
        coverage[label] = {"value": value, "covered_intervals": covered, "n": n,
                           "scope": "executed_first_day_scenarios"}
        rows.append({"metric": label, "value": value, "unit": "fraction", "n": n})
    payload["coverage"] = coverage
    return pd.DataFrame(rows), payload


def solver_statistics(daily):
    runtime = daily.runtime_seconds.to_numpy(dtype=float)
    achieved = daily.achieved_mip_gap.to_numpy(dtype=float)
    requested = daily.requested_mip_gap.to_numpy(dtype=float)
    if not (np.isfinite(runtime).all() and np.isfinite(achieved).all()
            and np.isfinite(requested).all()):
        raise ValueError("solver逐日统计包含非有限数值")
    return {"total_runtime": float(runtime.sum()),
            "mean_daily_runtime": float(runtime.mean()),
            "median_daily_runtime": float(np.median(runtime)),
            "max_daily_runtime": float(runtime.max()),
            "requested_mip_gap": float(requested.max()),
            "achieved_mip_gap_mean": float(achieved.mean()),
            "achieved_mip_gap_max": float(achieved.max()),
            "days_completed": len(daily)}


def write_final_artifacts(frame, actual, output):
    output = Path(output)
    validate_dispatch(frame, full_period=True)
    validation = final_validation(frame)
    write_json(output/"final_validation.json", validation)
    if not validation["overall_pass"]:
        raise ValueError("最终物理一致性验收失败")
    calibration = pd.read_csv(output/"calibration.csv", float_precision="round_trip")
    calibration_table, calibration_payload = final_calibration(frame, actual, calibration)
    write_csv(output/"final_calibration.csv", calibration_table)
    write_json(output/"final_calibration.json", calibration_payload)
    daily = pd.read_csv(output/"daily.csv", float_precision="round_trip")
    solver = solver_statistics(daily)
    write_json(output/"solver_statistics.json", solver)
    write_json(output/"summary.json", {
        **summary(frame),
        "model_configuration": {
            "tail_stratification": PROTOCOL["tail_stratification"],
            "tail_fraction": PROTOCOL["tail_fraction"],
            "dynamic_soc_reserve": PROTOCOL["dynamic_soc_reserve"],
            "reserve_quantile": PROTOCOL["reserve_quantile"],
            "evaluation_initial_soc_kwh": PROTOCOL["evaluation_initial_soc_kwh"]},
        "solver_statistics": solver})


def export_result(frame, template, output):
    validate_dispatch(frame, full_period=True)
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(template)
    expected = ["计划购电量", "充放电量", "紧急购电量"]
    if workbook.sheetnames != expected:
        raise ValueError("result2模板工作表结构不符")
    plans, storage, emergencies = [workbook[name] for name in expected]
    original_headers = [cell.value for cell in plans[1]]
    styles = {sheet.title: [copy(c._style) for c in sheet[2]] for sheet in workbook}
    for sheet in workbook:
        sheet.delete_rows(2, sheet.max_row)
    intervals = calendar_for_day("2025-02-01").interval.tolist()
    for col, label in enumerate(intervals, 2):
        plans.cell(1, col, label)
    plans.cell(1, 2).comment = Comment("原模板144列整体偏移10分钟；按方案2.2.2统一为00:00—24:00的144个区间。原标签已保存至template_audit.json。", "Q2模型")
    for row, (date, day) in enumerate(frame.groupby("date", sort=False), 2):
        plans.cell(row, 1, pd.Timestamp(date).to_pydatetime())
        for col, value in enumerate(day.grid, 2):
            plans.cell(row, col, float(value))
        # 模板为冻结仿真结果的数值输入区，费用由优化/回放引擎计算，非可编辑Excel计算模型。
        plans.cell(row, 146, float(day.grid.sum()))
        plans.cell(row, 147, float(day.planned_cost.sum()))
    blocks = aggregate_storage(frame)
    for i, item in enumerate(blocks.itertuples(), 2):
        block = (i-2)%6
        for col, value in enumerate((pd.Timestamp(item.date).to_pydatetime() if block == 0 else None,
                                     item.interval, item.charge_kwh, item.discharge_kwh,
                                     "0:00" if block == 0 else "24:00" if block == 1 else None,
                                     item.soc_at_0000 if block == 0 else item.soc_at_2400 if block == 1 else None), 1):
            storage.cell(i, col, value)
    for i, item in enumerate(emergency_events(frame).itertuples(), 2):
        emergencies.cell(i, 1, pd.Timestamp(item.date).to_pydatetime())
        emergencies.cell(i, 2, item.interval)
        emergencies.cell(i, 3, item.emergency_kwh)
    for sheet in workbook:
        sheet.freeze_panes = "B2" if sheet == plans else "A2"
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell._style = copy(styles[sheet.title][cell.column-1])
                if isinstance(cell.value, (float, int)):
                    cell.number_format = "0.000000"
    pending = output/"result2.pending.xlsx"
    workbook.save(pending); workbook.close()
    check = load_workbook(pending, read_only=True, data_only=True)
    try:
        saved = np.asarray([row[1:145] for row in list(check["计划购电量"].values)[1:]], dtype=float)
        np.testing.assert_allclose(saved.ravel(), frame.grid, atol=1e-9, rtol=1e-12)
        totals = np.asarray([row[145:147] for row in list(check["计划购电量"].values)[1:]], dtype=float)
        np.testing.assert_allclose(totals, frame.groupby("date", sort=False)[["grid", "planned_cost"]].sum(), atol=1e-8)
        if check["充放电量"].max_row != 1+334*6:
            raise ValueError("充放电表不是334天×6个区间")
        for i, row in enumerate(list(check["充放电量"].values)[1:]):
            np.testing.assert_allclose(row[2:4], blocks.iloc[i][["charge_kwh", "discharge_kwh"]].to_numpy(dtype=float), atol=1e-9)
            if i%6 in (0, 1):
                expected_soc = blocks.iloc[i]["soc_at_0000" if i%6 == 0 else "soc_at_2400"]
                np.testing.assert_allclose(row[5], expected_soc, atol=1e-9)
        total_emergency = sum(row[2] for row in list(check["紧急购电量"].values)[1:])
        if not np.isclose(total_emergency, frame.emergency.sum(), rtol=1e-12, atol=1e-8):
            raise ValueError("紧急区间汇总不守恒")
    finally:
        check.close()
    pending.replace(output/"result2.xlsx")
    write_json(output/"template_audit.json", {"original_plan_headers": original_headers,
                                             "corrected_interval_headers": intervals,
                                             "reason": "template_shift_conflicts_with_144_source_day_intervals",
                                             "full_period_rows": len(frame), "numeric_roundtrip": True})
    write_tables(frame, output)
