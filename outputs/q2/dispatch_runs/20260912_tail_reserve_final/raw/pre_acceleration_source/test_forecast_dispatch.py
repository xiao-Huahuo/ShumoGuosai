"""新版预测与外部输出的独立验收；所有合成数据只存在于测试临时目录。"""

import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from comparators import deterministic_policy, oracle_pair, rigid_schedule
from data import ROOT
from export_dispatch import (emergency_events, export_result, final_calibration,
                             final_validation, solver_statistics, validate_dispatch)
from forecast import economic_metrics, harmonic_forecast, lightgbm_forecast, load_features, seasonal, select_pipeline
from optimization import solve_policy
from policy import CAP, E_MIN, Policy, replay
from protocol import PROTOCOL
from relaxation import recourse_bound
from rolling import (ComputationalLimit, _exclusive_run_lock, dispatch_frame, make_day,
                     run_rollout, warm_start)
from shadows import freeze_lightgbm
from incumbent import value_gradient
from checkpoint import read_checkpoint, recover_checkpoint, save_checkpoint
from dispatch import (forecast_hashes, load_prepared, prepared_hashes,
                      validate_prepared_protocol)
from data import write_json
from protocol import signature
from experiments import load_richer_policy, richer_context, save_richer_policy
from final_report import visual_review_current
from data import sha256


def history_fixture(days=14):
    slot = np.arange(144)
    history = np.zeros((days, 144, 2))
    for d in range(days):
        history[d, :, 0] = 1000+100*(d%7)+200*np.cos(2*np.pi*slot/144)
        history[d, :, 1] = (300+d)*np.maximum(np.sin(2*np.pi*(slot-36)/144), 0)
    return history


class RicherCheckpointTests(unittest.TestCase):
    def fixture(self, output):
        audit, plan = output/"reference.json", output/"reference.csv"
        audit.write_text("{}", encoding="utf-8"); plan.write_text("reference", encoding="utf-8")
        net, weights, prices = np.zeros((2, 144)), np.array([.5, .5]), np.ones(144)
        policy = Policy(np.zeros(144), np.tile([100., 200., 300.], (144, 1)),
                        np.tile([300., 200., 100.], (144, 1)))
        context = richer_context(net, weights, prices, 6000, audit, plan)
        record = {"date": "2025-03-20", "K": 1, "S": 2,
                  "solver": {"reliable": True, "objective": 0., "dual_bound": 0.}}
        return policy, record, context, net, prices, weights

    def test_richer_policy_roundtrip_replays_and_checks_global_certificate(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            policy, record, context, net, prices, weights = self.fixture(output)
            save_richer_policy(output, record, policy, context)
            recovered, saved = load_richer_policy(output, record["date"], context, net, prices, weights, 6000)
            np.testing.assert_array_equal(recovered.charge_cap, policy.charge_cap)
            np.testing.assert_array_equal(recovered.discharge_cap, policy.discharge_cap)
            self.assertEqual(saved["solver"]["objective"], 0.)
            saved["solver"]["objective"] = 10.
            write_json(output/f"richer_soc/{record['date']}.json", saved)
            with self.assertRaises(AssertionError):
                load_richer_policy(output, record["date"], context, net, prices, weights, 6000)

    def test_richer_cache_rejects_wrong_scenarios_and_changed_policy_bytes(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            policy, record, context, net, prices, weights = self.fixture(output)
            saved = save_richer_policy(output, record, policy, context)
            with self.assertRaisesRegex(ValueError, "不一致"):
                load_richer_policy(output, record["date"], {**context, "scenario_input_sha256": "changed"}, net, prices, weights, 6000)
            path = output/saved["policy_file"]
            path.write_text(path.read_text(encoding="utf-8")+"\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "摘要不符"):
                load_richer_policy(output, record["date"], context, net, prices, weights, 6000)

    def test_richer_interrupted_new_receipt_preserves_previous_complete_policy(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            policy, record, context, net, prices, weights = self.fixture(output)
            original = save_richer_policy(output, record, policy, context)
            with patch("experiments.atomic_json", side_effect=OSError("simulated interrupted commit")):
                with self.assertRaises(OSError):
                    save_richer_policy(output, record, policy, context)
            _, restored = load_richer_policy(output, record["date"], context, net, prices, weights, 6000)
            self.assertEqual(restored["policy_file"], original["policy_file"])


class ArtifactAcceptanceTests(unittest.TestCase):
    def test_regenerated_plot_cannot_inherit_previous_visual_approval(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp)
            (output/"figures").mkdir()
            png = output/"figures/real_rollout.png"
            png.write_bytes(b"synthetic previously reviewed image fixture")
            write_json(output/"figures/manifest.json", {"figures": ["real_rollout"]})
            write_json(output/"figures/visual_acceptance.json", {
                "source_sha256": sha256(ROOT/"src/plots/q2_dispatch.py"),
                "reviewed_png_sha256": {"real_rollout": sha256(png)}})
            self.assertTrue(visual_review_current(output))
            png.write_bytes(b"synthetic regenerated different image fixture")
            self.assertFalse(visual_review_current(output))


class ForecastTests(unittest.TestCase):
    def test_recursive_seasonal_uses_its_own_previous_prediction(self):
        history = history_fixture()
        predicted = seasonal(history, 3)
        np.testing.assert_array_equal(predicted[0, :, 0], history[-7, :, 0])
        np.testing.assert_allclose(predicted[1, :, 1], np.mean(np.r_[history[-6:, :, 1], predicted[0:1, :, 1]], axis=0))
        np.testing.assert_allclose(predicted[2, :, 1], np.mean(np.r_[history[-5:, :, 1], predicted[:2, :, 1]], axis=0))

    def test_exact_eight_load_features_and_actual_candidate(self):
        history = history_fixture()
        features = load_features(history[:, :, 0], "2025-01-15")
        self.assertEqual(features.shape, (144, 8))
        np.testing.assert_array_equal(features[:, 3:6], history[[-1, -2, -7], :, 0].T)
        expected = history.copy()
        prediction = lightgbm_forecast(history, "2025-01-15", 3, PROTOCOL["lightgbm_candidates"][0])
        self.assertEqual(prediction.shape, (3, 144))
        self.assertTrue(np.isfinite(prediction).all())
        np.testing.assert_array_equal(history, expected)

    def test_dhr_orders_and_diagnostics_use_fixed_dictionary(self):
        prediction, diagnostics, correlation = harmonic_forecast(history_fixture(), 3)
        self.assertEqual(prediction.shape, (3, 144))
        self.assertTrue((prediction >= 0).all())
        self.assertFalse(diagnostics["annual_harmonics"])
        self.assertEqual(len(diagnostics["candidates"]), 8)
        self.assertEqual(len(correlation), 145)
        self.assertEqual([v["lag"] for v in diagnostics["ljung_box"]], [6, 144])

    def test_freezing_cannot_use_february_hyperparameter_performance(self):
        actual = np.zeros((40, 144, 2))
        store = {"load_lgb_0": np.zeros((40, 3, 144)), "load_lgb_1": np.ones((40, 3, 144))}
        store["load_lgb_0"][31:] = 1e9
        self.assertEqual(freeze_lightgbm(store, actual)[0], 0)
        actual[31:] = 1e12
        self.assertEqual(freeze_lightgbm(store, actual)[0], 0)

    def test_january_storage_is_inactive_and_february_soc_is_6000(self):
        actual = np.zeros((365, 144, 2))
        dates = pd.date_range("2025-01-01", periods=365)
        with tempfile.TemporaryDirectory() as directory:
            initial = warm_start(actual, dates, np.ones(144), Path(directory))
            record = json.loads((Path(directory)/"warm_start.json").read_text(encoding="utf-8"))
            self.assertEqual(initial, 6000.)
            self.assertEqual(record["feb1_baseline_soc"], 6000.)
            self.assertEqual(record["january_charge_kwh"], 0.)
            self.assertEqual(record["january_discharge_kwh"], 0.)

    def test_four_to_one_loss_and_lexicographic_selection(self):
        actual = np.array([[[6., 0.], [6., 0.]]])
        prediction = np.array([[[0., 0.], [12., 0.]]])
        result = economic_metrics(actual, prediction, np.ones(2))
        self.assertEqual(result["econ"], 2.5)
        scores = {"baseline": dict(result), "complex": {**result, "econ": 3., "net_mae": 0}}
        self.assertEqual(select_pipeline(scores), "baseline")


class ComparatorTests(unittest.TestCase):
    def test_incumbent_gradient_matches_finite_differences_away_from_kinks(self):
        net = np.array([[100., 400, 450, 100], [90., 350, 400, 120]])
        vector = np.r_[np.full(4, 200.), np.full(4, 80.), np.full(4, 30.)]
        prices, weights = np.array([.4, 1, .8, .5]), np.array([.5, .5])
        _, gradient = value_gradient(vector, net, prices, weights, 5000)
        finite = [(value_gradient(vector+step*1e-4, net, prices, weights, 5000)[0]-
                   value_gradient(vector-step*1e-4, net, prices, weights, 5000)[0])/2e-4 for step in np.eye(len(vector))]
        np.testing.assert_allclose(gradient, finite, atol=1e-6, rtol=1e-6)

    def test_richer_gradient_uses_actual_half_open_soc_band(self):
        net = np.array([[700., -900, 1000, -500], [600., -600, 900, -400]])
        grid = np.full(4, 200.)
        cp = np.tile([120., 400., 500.], (4, 1))
        rp = np.tile([180., 450., 600.], (4, 1))
        vector = np.r_[grid, cp.ravel(), rp.ravel()]
        prices, weights = np.array([.4, 1, .8, .5]), np.array([.5, .5])
        value, gradient = value_gradient(vector, net, prices, weights, 4480)
        exact = prices@grid+sum(5*p*(prices@replay(Policy(grid, cp, rp), path, 4480)["emergency"]) for p, path in zip(weights, net))
        self.assertAlmostEqual(value, exact)
        finite = [(value_gradient(vector+step*1e-4, net, prices, weights, 4480)[0]-
                   value_gradient(vector-step*1e-4, net, prices, weights, 4480)[0])/2e-4 for step in np.eye(len(vector))]
        np.testing.assert_allclose(gradient, finite, atol=1e-5, rtol=1e-5)

    def test_deterministic_certificate_matches_exact_model(self):
        net, prices = np.array([1000., -900, 700, 200]), np.array([1., .4, 1., .4])
        _, _, certificate = deterministic_policy(net, prices, E_MIN)
        _, _, exact = solve_policy(net[None], prices, np.array([1.]), E_MIN, time_limit=15, gap=1e-8)
        self.assertTrue(certificate["reliable"])
        self.assertAlmostEqual(certificate["objective"], exact["objective"], places=5)

    def test_lower_bound_is_not_executed_and_oracles_are_nested(self):
        net = np.array([[1000., -900, 700, 200], [1300, -400, 900, 100]])
        prices, weights = np.array([1., .4, 1., .4]), np.array([.4, .6])
        policy, responses, bound = recourse_bound(net, prices, weights, E_MIN)
        _, _, exact = solve_policy(net, prices, weights, E_MIN, time_limit=15, gap=1e-8)
        self.assertLessEqual(bound["lower_bound"], exact["objective"]+1e-5)
        self.assertLessEqual(exact["objective"], bound["incumbent_cost"]+1e-5)
        real_cost = float(prices@(policy.grid+5*responses[0]["emergency"]))
        oracle = oracle_pair(net[0], prices, E_MIN, policy.grid, real_cost)
        self.assertGreaterEqual(oracle["Gap_resp"], -1e-8)

    def test_richer_policy_contains_main_policy_at_soc_boundaries(self):
        base = Policy([20, 0, 0], [300]*3, [200]*3)
        expanded = Policy(base.grid, np.repeat(base.charge_cap[:, None], 3, axis=1),
                          np.repeat(base.discharge_cap[:, None], 3, axis=1))
        net = np.array([[300., -200, 700], [200, -300, 800]])
        _, responses, info = solve_policy(net, np.ones(3), [.5, .5], 4400,
                                          segments=3, fixed_policy=expanded, time_limit=15)
        self.assertTrue(info["reliable"])
        for path, actual in zip(net, responses):
            expected = replay(base, path, 4400)
            for key in expected:
                np.testing.assert_allclose(actual[key], expected[key], atol=1e-5)


class ExportTests(unittest.TestCase):
    def test_final_validation_calibration_and_solver_statistics(self):
        policy = Policy(np.zeros(144), np.zeros(144), np.zeros(144), np.full(144, 100.))
        frame = dispatch_frame("2025-02-01", policy, np.zeros(144), 6000, np.ones(144))
        frame["predicted_load_kw"] = 0.
        frame["predicted_pv_kw"] = 0.
        validation = final_validation(frame)
        self.assertTrue(validation["overall_pass"])
        self.assertEqual(validation["reserve_discharge_violation_count"], 0)
        actual = np.zeros((365, 144, 2))
        coverage = pd.DataFrame([
            {"horizon": 0, "nominal": .8, "covered_intervals": 120, "n": 144},
            {"horizon": 0, "nominal": .9, "covered_intervals": 130, "n": 144}])
        table, payload = final_calibration(frame, actual, coverage)
        self.assertEqual(len(table), 8)
        self.assertAlmostEqual(payload["coverage"]["PICP80"]["value"], 120/144)
        daily = pd.DataFrame({"runtime_seconds": [2., 4.], "achieved_mip_gap": [.01, .02],
                              "requested_mip_gap": [.03, .03]})
        stats = solver_statistics(daily)
        self.assertEqual(stats["total_runtime"], 6.)
        self.assertEqual(stats["days_completed"], 2)

    def test_cross_midnight_events_and_empty_events(self):
        times = pd.date_range("2025-02-01 23:50", periods=4, freq="10min")
        frame = pd.DataFrame({"timestamp": times, "emergency": [0., 1., 2., 0.], "emergency_cost": [0., 5., 10., 0.]})
        events = emergency_events(frame)
        self.assertEqual(len(events), 1)
        self.assertEqual(events.emergency_kwh.iloc[0], 3)
        self.assertEqual(events.interval.iloc[0], "23:50-00:10+1日")
        frame["emergency"] = 0
        self.assertTrue(emergency_events(frame).empty)

    def test_complete_template_roundtrip_and_partial_rejection(self):
        dates = pd.date_range("2025-02-01", "2025-12-31")
        # 全期合成零负荷用来核验布局/日期/单位，绝不写入正式运行目录。
        policy = Policy(np.zeros(144), np.zeros(144), np.zeros(144))
        frame = pd.concat([dispatch_frame(d, policy, np.zeros(144), 6000, np.ones(144)) for d in dates], ignore_index=True)
        with self.assertRaises(ValueError):
            validate_dispatch(frame.iloc[:144], full_period=True)
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            export_result(frame, ROOT/"docs/CUMCM2026Problems/C题/附件/附件5/result2.xlsx", output)
            wb = load_workbook(output/"result2.xlsx", read_only=True, data_only=True)
            self.assertEqual(wb["计划购电量"].max_row, 335)
            self.assertEqual(wb["计划购电量"]["B1"].value, "00:00-00:10")
            self.assertEqual(wb["计划购电量"].cell(1, 145).value, "23:50-24:00")
            self.assertEqual(wb["充放电量"].max_row, 2005)
            wb.close()
            self.assertTrue(json.loads((output/"template_audit.json").read_text(encoding="utf-8"))["numeric_roundtrip"])


class RollingTests(unittest.TestCase):
    def test_interrupted_csv_commit_recovers_complete_old_or_new_checkpoint(self):
        for failed_file in ("dispatch.csv", "daily.csv", "calibration.csv", "checkpoint.json"):
            with self.subTest(failed_file=failed_file), tempfile.TemporaryDirectory() as directory:
                output = Path(directory)
                first = pd.DataFrame({"value": [1]})
                second = pd.DataFrame({"value": [1, 2]})
                save_checkpoint(output, first, [{"day": 1}], [{"score": 1}])
                replace = Path.replace
                def interrupted(path, target):
                    result = replace(path, target)
                    if target == output/failed_file:
                        raise InterruptedError("模拟在文件替换后进程中断")
                    return result
                with patch.object(Path, "replace", interrupted), self.assertRaises(InterruptedError):
                    save_checkpoint(output, second, [{"day": 1}, {"day": 2}], [{"score": 1}, {"score": 2}])
                recover_checkpoint(output)
                expected = 2 if failed_file == "checkpoint.json" else 1
                manifest = json.loads((output/"checkpoint.json").read_text(encoding="utf-8"))
                self.assertEqual(manifest["days"], expected)
                for name in ("dispatch.csv", "daily.csv", "calibration.csv"):
                    self.assertEqual(len(pd.read_csv(output/name)), expected)
                self.assertFalse((output/"checkpoint_transaction.json").exists())

    def test_interrupted_first_day_leaves_no_false_checkpoint(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            replace = Path.replace
            def interrupted(path, target):
                result = replace(path, target)
                if target == output/"daily.csv":
                    raise InterruptedError("首次提交中断")
                return result
            with patch.object(Path, "replace", interrupted), self.assertRaises(InterruptedError):
                save_checkpoint(output, pd.DataFrame({"value": [1]}), [{"day": 1}], [{"score": 1}])
            recover_checkpoint(output)
            self.assertFalse((output/"dispatch.csv").exists())
            self.assertFalse((output/"checkpoint.json").exists())

    def test_supplementary_compute_only_after_entire_day_failed(self):
        actual = np.zeros((33, 144, 2)); dates = pd.date_range("2025-01-01", periods=33)
        seen = []
        def solver(store, frozen, history, d, prices, soc, **settings):
            budget = settings.get("solver_seconds", 120)
            seen.append((d, len(history), budget))
            if d == 32 and budget == 120:
                raise ComputationalLimit({"original_failed": True})
            return Policy(np.zeros(144), np.zeros(144), np.zeros(144)), np.ones(1), np.zeros((1, 144)), np.zeros((1, 144, 2)), {
                "selected_K": 1, "computational_limited": False,
                "1": {"M": 14, "selected_solver": {"reliable": True}}}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            with patch("rolling.make_day", side_effect=solver):
                result, _ = run_rollout({}, 0, actual, dates, np.ones(144), 6000, output, rescue_seconds=900)
            self.assertEqual(seen, [(31, 31, 120), (32, 32, 120), (32, 32, 900)])
            self.assertEqual(result.date.nunique(), 2)
            audit = json.loads((output/"daily_audit/2025-02-02.json").read_text(encoding="utf-8"))
            self.assertTrue(audit["computational_limited"])
            self.assertEqual(audit["original_budget_failure"], {"original_failed": True})
            self.assertTrue(list((output/"failed_attempts").glob("*.json")))

    def test_failed_supplement_does_not_skip_date_or_erase_previous_day(self):
        actual = np.zeros((34, 144, 2)); dates = pd.date_range("2025-01-01", periods=34)
        seen = []
        def solver(store, frozen, history, d, prices, soc, **settings):
            seen.append(d)
            if d != 31:
                raise ComputationalLimit({"all_candidates_failed": True})
            return Policy(np.zeros(144), np.zeros(144), np.zeros(144)), np.ones(1), np.zeros((1, 144)), np.zeros((1, 144, 2)), {
                "selected_K": 1, "computational_limited": False,
                "1": {"M": 14, "selected_solver": {"reliable": True}}}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            with patch("rolling.make_day", side_effect=solver), self.assertRaises(ComputationalLimit):
                run_rollout({}, 0, actual, dates, np.ones(144), 6000, output, rescue_seconds=900)
            self.assertEqual(seen, [31, 32, 32])
            self.assertEqual(len(pd.read_csv(output/"dispatch.csv")), 144)
            self.assertEqual(json.loads((output/"checkpoint.json").read_text(encoding="utf-8"))["days"], 1)

    def test_same_experiment_cannot_have_two_writers(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            with _exclusive_run_lock(output/"run.lock"):
                with self.assertRaisesRegex(RuntimeError, "已有写入进程"):
                    run_rollout({}, 0, None, None, None, None, output)

    def test_prepared_generated_data_tampering_is_rejected_before_loading(self):
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            write_json(output/"raw/protocol.json", {"protocol_sha256": signature(), "forecast_source_sha256": forecast_hashes()})
            write_json(output/"inputs/processed/source_manifest.json", {})
            data = output/"inputs/processed/timeseries.csv"
            data.write_text("load_kw,pv_kw\n1,2\n", encoding="utf-8")
            write_json(output/"raw/prepared_manifest.json", {"sha256": prepared_hashes(output)})
            data.write_text("load_kw,pv_kw\n9,2\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "摘要改变"):
                load_prepared(output)

    def test_prepared_forecasts_remain_reusable_when_only_dispatch_model_changes(self):
        old = {"protocol_sha256": "old-model-signature",
               "forecast_source_sha256": {**forecast_hashes(), "scenarios.py": "old",
                                           "policy.py": "old", "protocol.py": "old"},
               **{key: PROTOCOL[key] for key in
                  ("history_boundary", "forecast_horizons", "shadow_start",
                   "lightgbm_candidates", "lightgbm_training_window",
                   "lightgbm_freeze_metric", "dhr_harmonics", "dhr_ar_orders",
                   "dhr_ma_orders", "dhr_history_days", "dhr_order_rule",
                   "dhr_max_iterations", "dhr_difference_order")}}
        validate_prepared_protocol(old)
        old["dhr_history_days"] += 1
        with self.assertRaisesRegex(ValueError, "预测冻结参数"):
            validate_prepared_protocol(old)

    def test_real_selection_metadata_serializes_numpy_origin_indices(self):
        actual = np.zeros((34, 144, 2))
        store = {key: np.zeros((34, 3, 144)) for key in
                 ("load_week", "load_lgb_0", "pv_mean7", "pv_dhr")}
        for values in store.values():
            values[:14] = np.nan
        _, _, _, _, audit = make_day(store, 0, actual[:31], 31, np.ones(144), 6000)
        serialized = json.loads(json.dumps(audit))
        self.assertEqual(serialized["selected_K"], 2)
        self.assertTrue(all(j < 31 for j in serialized["2"]["selection"]["common_scored_origins"]))

    def test_resume_replays_frozen_policy_and_rejects_corrupted_checkpoint(self):
        actual = np.zeros((33, 144, 2)); dates = pd.date_range("2025-01-01", periods=33)
        seen = []
        def day_solver(store, frozen, history, d, prices, soc, **settings):
            seen.append((d, len(history)))
            policy = Policy(np.zeros(144), np.zeros(144), np.zeros(144), np.full(144, 100.))
            return policy, np.ones(1), np.zeros((1, 144)), np.zeros((1, 144, 2)), {
                "selected_K": 1, "computational_limited": False,
                "1": {"M": 14, "selected_solver": {"reliable": True}}}
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory)
            with patch("rolling.make_day", side_effect=day_solver):
                first, _ = run_rollout({}, 0, actual, dates, np.ones(144), 6000, output)
            self.assertEqual(seen, [(31, 31), (32, 32)])
            with patch("rolling.make_day", side_effect=AssertionError("已完成前缀不得重新优化")):
                resumed, _ = run_rollout({}, 0, actual, dates, np.ones(144), 6000, output)
            np.testing.assert_array_equal(resumed.grid, first.grid)
            np.testing.assert_array_equal(resumed.soc_reserve_kwh, first.soc_reserve_kwh)
            changed = pd.read_csv(output/"dispatch.csv")
            changed.loc[144:, "initial_soc"] = 7000
            changed.to_csv(output/"dispatch.csv", index=False, encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "摘要不符"):
                read_checkpoint(output, actual, dates, np.ones(144), 6000, 31, 33, {})
            (output/"checkpoint.json").unlink()
            with self.assertRaisesRegex(ValueError, "跨日SOC"):
                read_checkpoint(output, actual, dates, np.ones(144), 6000, 31, 33, {})

    def test_execution_cannot_call_optimizer_and_prefix_cannot_see_future(self):
        policy = Policy(np.zeros(144), np.full(144, 100.), np.full(144, 100.))
        net = np.zeros(144); net[:2] = [200, -200]
        with patch("rolling._solve", side_effect=AssertionError("日内不能求解")):
            first = dispatch_frame("2025-02-01", policy, net, 6000, np.ones(144))
            net[2:] = 1e9
            second = dispatch_frame("2025-02-01", policy, net, 6000, np.ones(144))
        pd.testing.assert_frame_equal(first.iloc[:2], second.iloc[:2])


if __name__ == "__main__":
    unittest.main()
