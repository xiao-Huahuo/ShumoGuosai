"""针对时间错位、约束遗漏及输出篡改的独立验收；合成数据仅用于测试。"""

import copy
import csv
import datetime as dt
import json
import tempfile
import unittest
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from export import check_csv_output, read_csv
from model import (PHYS_TOL_KWH, COST_TOL_YUAN, ARTIFACT_ROOT, ROOT, SCENARIOS, arrays, check_solution, endpoint_minutes,
                   read_inputs, solve, write_csv)


class Q1Acceptance(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.input_path = ARTIFACT_ROOT / "inputs/raw/attachment1.xlsx"
        cls.data = read_inputs(cls.input_path)
        cls.solution = solve(cls.data, .9)  # 当前代码重求，CSV必须与本轮结果对应。
        cls.output = ARTIFACT_ROOT / "processed"

    def test_all_actual_scenarios(self):
        for key, efficiency in SCENARIOS.items():
            with self.subTest(scenario=key):
                s = json.loads((ARTIFACT_ROOT / f"raw/solution_{key}.json").read_text(encoding="utf-8"))
                self.assertEqual(s["eta_c"], efficiency)
                self.assertEqual(s["eta_d"], efficiency)
                self.assertTrue(check_solution(self.data, s)["passed"])

    def test_time_and_unit_examples(self):
        self.assertEqual(self.data[0]["interval"], "00:00-00:10")
        self.assertEqual(self.data[-1]["interval"], "23:50-24:00")
        self.assertEqual(self.data[60]["interval"], "10:00-10:10")
        self.assertEqual(endpoint_minutes(self.data[60]["source_endpoint"]), 610)
        self.assertAlmostEqual(self.data[0]["load_kwh"], 3439.8466 / 6, places=10)
        self.assertEqual(len({r["interval"] for r in self.data}), 144)
        for value in (dt.time(0, 10), dt.timedelta(minutes=10), 10/1440, "00:10"):
            self.assertEqual(endpoint_minutes(value), 10)
        for value in (1.0, dt.timedelta(days=1), "0:00+1", "24:00"):
            self.assertEqual(endpoint_minutes(value), 1440)

    def test_corrupted_inputs_rejected(self):
        mutations = {"shifted_endpoint": ("A2", "00:20"), "duplicate_endpoint": ("A3", "00:10"),
                     "missing_load": ("C2", None), "negative_pv": ("D2", -1),
                     "negative_price": ("B2", -1), "numeric_text": ("B2", "0.4248"),
                     "invalid_minute": ("A2", "00:70")}
        with tempfile.TemporaryDirectory() as tmp:
            for name, (cell, value) in mutations.items():
                with self.subTest(case=name):
                    wb = load_workbook(self.input_path)
                    wb.active[cell] = value
                    path = Path(tmp) / f"{name}.xlsx"
                    wb.save(path)
                    wb.close()
                    with self.assertRaises(ValueError):
                        read_inputs(path)
            wb = load_workbook(self.input_path)
            wb.active.delete_rows(145)
            path = Path(tmp) / "missing_row.xlsx"
            wb.save(path)
            wb.close()
            with self.assertRaises(ValueError):
                read_inputs(path)

    def test_corrupted_schedules_rejected(self):
        cases = [("G", 0, -1), ("C", 0, 1000), ("D", 0, 1), ("E", 0, 12001),
                 ("W", 0, 1), ("u", 0, 0.5), ("E", 143, 5999), ("G", 10, float("nan"))]
        for name, index, value in cases:
            with self.subTest(variable=name, index=index):
                s = copy.deepcopy(self.solution)
                s[name][index] = value
                with self.assertRaises(ValueError):
                    check_solution(self.data, s)
        for key, value in (("objective_yuan", 0), ("dual_bound_yuan", 0), ("status", 1),
                           ("mip_gap", float("nan"))):
            with self.subTest(solver_field=key):
                s = copy.deepcopy(self.solution)
                s["solver"][key] = value
                with self.assertRaises(ValueError):
                    check_solution(self.data, s)
        s = copy.deepcopy(self.solution)
        s["E0"] = 5000
        with self.assertRaises(ValueError):
            check_solution(self.data, s)

    def test_csv_roundtrip_and_tamper_rejection(self):
        self.assertTrue(check_csv_output(self.output, self.data, self.solution)["passed"])
        cases = [("result1.csv", 0, "interval", "00:10-00:20"),
                 ("result1.csv", 0, "G_kwh", -1), ("result1.csv", 0, "C_kwh", 1000),
                 ("result1.csv", 0, "G_kwh", "nan"), ("result1.csv", 0, "E_start_kwh", 5999),
                 ("result1.csv", 2, "soc_fraction", 0), ("result1.csv", 0, "cost_yuan", 0),
                 ("table2_storage.csv", 0, "charge_kwh", 1),
                 ("daily_summary.csv", 0, "terminal_kwh", 5999),
                 ("daily_summary.csv", 0, "load_kwh", "nan"),
                 ("result1.csv", 0, "cost_yuan", "nan"),
                 ("table1_purchase.csv", 0, "grid_kwh", 1)]
        cases.append(("table1_purchase.csv", 3, "grid_kwh", "nan"))
        files = ["result1.csv", "table1_purchase.csv", "table2_storage.csv", "daily_summary.csv"]
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            for file, index, field, value in cases:
                with self.subTest(file=file, field=field):
                    for name in files:
                        write_csv(directory / name, read_csv(self.output / name))
                    rows = read_csv(directory / file)
                    rows[index][field] = value
                    write_csv(directory / file, rows)
                    with self.assertRaises(ValueError):
                        check_csv_output(directory, self.data, self.solution)
            # 保留总和，交换两个时段的购电量，也必须被发现。
            for name in files:
                write_csv(directory / name, read_csv(self.output / name))
            rows = read_csv(directory / "result1.csv")
            rows[0]["G_kwh"], rows[18]["G_kwh"] = rows[18]["G_kwh"], rows[0]["G_kwh"]
            write_csv(directory / "result1.csv", rows)
            with self.assertRaises(ValueError):
                check_csv_output(directory, self.data, self.solution)

    def test_analytical_arbitrage_case(self):
        # 专用合成测试：首段0.1元，其余1元；末段负荷100kWh，无光伏。
        # 最优需首段购电100/(0.9²)，末段放电100；准确费用1000/81。
        synthetic = [{"price_yuan_per_kwh": 0.1 if t == 0 else 1.0,
                      "load_kwh": 100.0 if t == 143 else 0.0, "pv_kwh": 0.0} for t in range(144)]
        s = solve(synthetic, 0.9)
        self.assertTrue(check_solution(synthetic, s)["passed"])
        self.assertAlmostEqual(s["solver"]["objective_yuan"], 1000 / 81, delta=PHYS_TOL_KWH)
        self.assertAlmostEqual(sum(s["C"]), 100 / 0.81, delta=PHYS_TOL_KWH)
        self.assertAlmostEqual(sum(s["D"]), 100, delta=PHYS_TOL_KWH)

    def test_curtailment_is_available_when_absorption_is_limited(self):
        # 专用压力测试：每段2000kWh光伏且零负荷。日周期+互斥禁止无限吸收。
        synthetic = [{"price_yuan_per_kwh": 1.0, "load_kwh": 0.0, "pv_kwh": 2000.0} for _ in range(144)]
        s = solve(synthetic, 0.9)
        self.assertTrue(check_solution(synthetic, s)["passed"])
        self.assertAlmostEqual(sum(s["G"]), 0, delta=PHYS_TOL_KWH)
        self.assertGreater(sum(s["W"]), 0)

    def test_export_csv_matches_all_variables_and_inputs(self):
        with (ARTIFACT_ROOT / "processed/result1.csv").open(encoding="utf-8", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 144)
        p, _, _ = arrays(self.data)
        for key in ("G", "C", "D", "E", "W"):
            np.testing.assert_allclose([float(r[f"{key}_kwh"]) for r in rows], self.solution[key], rtol=0, atol=PHYS_TOL_KWH)
        np.testing.assert_allclose([float(r["u"]) for r in rows], self.solution["u"], rtol=0, atol=PHYS_TOL_KWH)
        self.assertAlmostEqual(sum(float(r["cost_yuan"]) for r in rows), p @ self.solution["G"], delta=PHYS_TOL_KWH)
        with (ARTIFACT_ROOT / "inputs/processed/timeseries.csv").open(encoding="utf-8", newline="") as stream:
            inputs = list(csv.DictReader(stream))
        self.assertEqual(len(inputs), 144)
        for i, record in enumerate(inputs):
            for key, value in self.data[i].items():
                self.assertEqual(record[key], str(value))


if __name__ == "__main__":
    unittest.main(verbosity=2)
