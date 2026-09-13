"""3.3—3.8：原始附件校验、Q2只读继承与严格可用时间边界。"""
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import sys
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from scipy.interpolate import PchipInterpolator
from .config import ROOT, ATTACHMENTS, NODES, DT, write_csv, write_json


@dataclass
class Inputs:
    actual: np.ndarray  # day × 144 × (load,pv), kW
    hourly: np.ndarray  # day × four origins × 24 horizons, kW
    prices: np.ndarray
    load_store: dict[str, np.ndarray]
    selected: list[str]
    provenance: dict

    def load(self, day: int, name: str | None = None) -> np.ndarray:
        if day < 7:
            raise ValueError('沿用Q2 1月1—7日无负荷预测、储能停机冷启动')
        name = name or self.selected[day]
        result = self.load_store[name][day, :2].reshape(-1).copy()
        if not np.isfinite(result).all():
            raise ValueError(f'Q2因果48h预测缺失：{day}/{name}')
        return result

    def observation(self, day: int, hour: int) -> float:
        endpoint = day*144+hour*6
        if endpoint == 0:
            raise ValueError('附件没有2025-01-01 00:00观测，不能把00:10真值当左端点')
        return float(self.actual.reshape(-1, 2)[endpoint-1, 1])

    def pv(self, day: int, hour: int, method: str = 'linear') -> np.ndarray:
        values = np.r_[self.observation(day, hour), self.hourly[day, hour//6]]
        points = np.arange(1, 145)/6
        if method == 'linear':
            return np.interp(points, np.arange(25), values)
        if method == 'pchip':
            return np.maximum(PchipInterpolator(np.arange(25), values)(points), 0)
        raise ValueError('未知插值方式')

    def history(self, day: int, hour: int, name: str, *, method: str = 'linear',
                forecast_hour: int | None = None, length: int = 144) -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
        """候选日j<d且所用轨迹终点<=当前节点；按Q2当前predictor取同名历史样本外残差。"""
        origin = hour if forecast_hour is None else forecast_hour
        offset = (hour-origin)*6
        if offset < 0 or offset+length > 144:
            raise ValueError('请求超过官方预报共同支持集')
        ids, load_errors, pv_errors, scales = [], [], [], []
        cutoff = day*144+hour*6
        flattened = self.actual.reshape(-1, 2)
        for j in range(7, day):
            start = j*144+hour*6
            if start+length > cutoff or start+length > len(flattened):
                continue
            forecast = self.load_store[name][j, :2].reshape(-1)[hour*6:hour*6+length]
            if not np.isfinite(forecast).all():
                continue
            pv_full = self.pv(j, origin, method)
            observed = flattened[start:start+length]
            ids.append(j)
            load_errors.append(observed[:, 0]-forecast)
            pv_errors.append(observed[:, 1]-pv_full[offset:offset+length])
            scales.append(pv_full.sum()*DT)
        return (np.asarray(ids, dtype=int), np.asarray(load_errors).reshape(-1, length),
                np.asarray(pv_errors).reshape(-1, length), np.asarray(scales))


def read_inputs(q2_run: Path, processed: Path | None = None) -> Inputs:
    # Q2模块采用脚本式导入；只读导入其已经验证的附件解析与季节预测函数。
    sys.path.insert(0, str(ROOT/'src/q2'))
    from data import read_attachment, endpoint_minutes
    from forecast import seasonal
    shadow = q2_run/'raw/shadows'
    meta = json.loads((shadow/'shape.json').read_text(encoding='utf-8'))
    if meta['boundary'] != 'source_date_before_origin' or tuple(meta['shape']) != (365, 3, 144):
        raise ValueError('Q2预测库来源/形状不符')
    folder = processed or ROOT/'inputs/q3/processed'
    frame, _ = read_attachment(ATTACHMENTS/'附件2.xlsx', folder/'actual_quality.csv')
    actual = frame[['load_kw', 'pv_kw']].to_numpy().reshape(365, 144, 2)
    wb = load_workbook(ATTACHMENTS/'附件1.xlsx', read_only=True, data_only=True)
    rows = list(wb.active.values)[1:]; wb.close()
    if [endpoint_minutes(r[0]) for r in rows] != list(range(10, 1441, 10)):
        raise ValueError('附件1电价右端点错位')
    prices = np.array([r[1] for r in rows], dtype=float)
    wb = load_workbook(ATTACHMENTS/'附件3.xlsx', read_only=True, data_only=True)
    rows = list(wb.active.values); wb.close()
    if len(rows) != 1461 or rows[0][2:] != tuple(f'预报{h}小时' for h in range(1, 25)):
        raise ValueError('附件3表头/形状不符')
    hourly = np.full((365, 4, 24), np.nan)
    day = None
    long_rows = []
    for row in rows[1:]:
        if row[0]:
            day = pd.Timestamp(row[0]).normalize()
        d = (day-pd.Timestamp('2025-01-01')).days
        hour = int(str(row[1]).split(':')[0])
        if hour not in NODES or not 0 <= d < 365 or np.isfinite(hourly[d, hour//6]).any():
            raise ValueError('附件3日期/发布时间重复或越界')
        hourly[d, hour//6] = np.array(row[2:], dtype=float)
        issue = day+pd.Timedelta(hours=hour)
        long_rows.extend({'issue': issue, 'target': issue+pd.Timedelta(hours=h),
                          'horizon_hours': h, 'forecast_kw': float(value)} for h, value in enumerate(row[2:], 1))
    if not np.isfinite(hourly).all() or (hourly < 0).any() or not np.isfinite(prices).all() or (prices <= 0).any():
        raise ValueError('原始预报或电价非法，不擅自填补')
    store = {name: np.array(np.memmap(shadow/f'{name}.dat', mode='r', dtype=meta['dtype'], shape=tuple(meta['shape'])))
             for name in meta['model_ids'] if name.startswith('load_')}
    # Q2冷启动只用季节负荷；保留其原始影子预测，并补齐7—13日同一因果流程。
    for d in range(7, 14):
        store['load_week'][d] = seasonal(actual[:d].copy(), 3)[..., 0]
    selected, decisions = ['unavailable']*7+['load_week']*24, []
    for d in range(31, 365):
        date = pd.Timestamp('2025-01-01')+pd.Timedelta(days=d)
        path = q2_run/f'processed/main/daily_audit/{date.date()}.json'
        audit = json.loads(path.read_text(encoding='utf-8'))
        name = audit[str(audit['selected_K'])]['selection']['pipeline'][0]
        selected.append(name)
        decisions.append({'date': str(date.date()), 'load_predictor': name, 'q2_audit': str(path),
                          'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    sources = [ATTACHMENTS/f'附件{i}.xlsx' for i in (1, 2, 3)]
    sources += [shadow/'shape.json', shadow/'frozen_lightgbm.json']+[shadow/f'{name}.dat' for name in store]
    provenance = {'q2_run': str(q2_run), 'q2_reselection': False, 'q2_decisions': decisions,
                  'files': {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in sources},
                  'cold_start': 'Jan1-Jan7 no load forecast/storage; Jan8-Jan31 seasonal; Q2 daily load identity thereafter',
                  'units': 'power kW; all optimization and replay energies kWh',
                  'right_endpoint': True, 'jan1_0000_pv': 'missing; never filled with future observation'}
    if processed is not None:
        write_csv(folder/'pv_forecasts.csv', pd.DataFrame(long_rows))
        write_csv(folder/'prices.csv', pd.DataFrame({'slot': np.arange(144), 'price': prices}))
        write_csv(folder/'q2_load_selection.csv', pd.DataFrame(decisions))
        write_json(folder/'source_manifest.json', provenance)
    return Inputs(actual, hourly, prices, store, selected, provenance)
